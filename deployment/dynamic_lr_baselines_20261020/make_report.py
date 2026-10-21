"""Write the final research report only after twelve verified endpoints.

Run with the bundled document Python runtime. The artifact-operation marker
and visual inspection of the rendered DOCX are separate required operations.
"""
import argparse
import json
from pathlib import Path

from gather_results import METHODS, ROOT, SCENES, sha


LABELS = {"cook_spinach": "Cook", "cut_roasted_beef": "Cut"}


def cost(index, method):
    resources = index.get("resources", {})
    seconds = index.get("training_seconds", resources.get("training_seconds"))
    if method == "wu":
        gpu, memory = "PRO6000 历史端点", None
    elif index.get("gpu_uuid") == "GPU-fd2728bf-7bb9-4ce0-fc1b-6ab34ae0a910":
        gpu, memory = "CTS RTX 3090", index.get("peak_memory_bytes")
    else:
        gpu = "A100 40GB"
        memory = index.get("peak_memory_bytes", resources.get("peak_torch_reserved_bytes"))
    decimal_gb = index.get("training_peak_GB") if method == "multi4d" else None
    gib = decimal_gb * 1e9 / 2**30 if decimal_gb is not None else (memory / 2**30 if memory else None)
    return gpu, "未知" if seconds is None else f"{seconds / 60:.2f}", "未知" if gib is None else f"{gib:.2f}"


def table_md(headers, rows):
    return ["| " + " | ".join(h.replace("\n", "<br>") for h in headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"] + [
        "| " + " | ".join(str(x).replace("\n", "<br>") for x in row) + " |" for row in rows]


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--comparison", type=Path, required=True)
    p.add_argument("--metadata", type=Path, required=True)
    p.add_argument("--visuals", type=Path, required=True)
    p.add_argument("--review", type=Path, required=True, help="Actual fixed-panel inspection notes and identities")
    p.add_argument("--out", type=Path, required=True, help="Output stem under docs")
    a = p.parse_args()
    comparison = json.loads(a.comparison.read_text())
    assert comparison["status"] == "complete" and comparison["completed"] == comparison["expected"] == 12
    records = {(r["method"], r["scene"]): r for r in comparison["records"]}
    assert set(records) == {(m, s) for m in METHODS for s in SCENES}
    metadata = {x["method"]: x for x in json.loads(a.metadata.read_text())["records"]}
    visual_receipt = json.loads((a.visuals / "visual_receipt.json").read_text())
    review = json.loads(a.review.read_text())
    assert visual_receipt["status"] == "completed" and len(visual_receipt["records"]) == 6
    assert review["status"] == "reviewed_by_agent" and review["human_acceptance"] is False
    inspected = {x["path"]: x["sha256"] for x in review["inspected_panels"]}
    for row in visual_receipt["records"]:
        for mode in ("whole", "face"):
            image = a.visuals / row[mode]["path"]
            assert sha(image) == row[mode]["sha256"] == inspected[str(image.resolve())]

    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.65), Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
        style = document.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Normal"].paragraph_format.line_spacing = Pt(14.5)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)
    document.styles["Title"].font.size = Pt(21)
    document.styles["Title"].paragraph_format.line_spacing = Pt(28)
    document.styles["Heading 1"].font.size = Pt(14)
    document.styles["Heading 1"].paragraph_format.line_spacing = Pt(20)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    document.styles["Caption"].font.size = Pt(9)
    document.styles["Caption"].font.bold = False
    document.styles["Caption"].paragraph_format.line_spacing = Pt(12)
    md = []

    def paragraph(text, style=None):
        document.add_paragraph(text, style)
        md.extend([text, ""])

    def heading(text):
        document.add_heading(text, 1)
        md.extend(["## " + text, ""])

    def picture(path, width):
        document.add_picture(str(path), width=Inches(width))
        para = document.paragraphs[-1]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing = 1
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.keep_with_next = True

    def table(headers, rows, widths, left_columns=(0,)):
        grid = document.add_table(rows=1, cols=len(headers))
        grid.autofit = False
        for column, width in zip(grid.columns, widths):
            column.width = Inches(width)
        props = grid._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            rule = OxmlElement("w:" + edge)
            for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                rule.set(qn("w:" + key), value)
            borders.append(rule)
        props.append(borders)
        for values in [headers, *rows]:
            cells = grid.rows[0].cells if values is headers else grid.add_row().cells
            cells[0]._tc.getparent().get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for i, (cell, value) in enumerate(zip(cells, values)):
                cell.width = Inches(widths[i])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = str(value)
                for para in cell.paragraphs:
                    para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(3)
                    para.paragraph_format.line_spacing = Pt(13)
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT if i in left_columns else WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.bold = values is headers
                if values is headers:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "E7EFF7")
                    cell._tc.get_or_add_tcPr().append(shading)
        repeat = OxmlElement("w:tblHeader")
        grid.rows[0]._tr.get_or_add_trPr().append(repeat)
        document.add_paragraph().paragraph_format.space_after = Pt(0)
        md.extend(table_md(headers, rows) + [""])

    title = "动态高斯原始低分辨率重建对比"
    document.add_paragraph(title, "Title")
    md.extend(["# " + title, ""])
    paragraph("已完成六种动态高斯重建方法在 Cook 和 Cut 两景的十二项固定端点评价。所有方法使用原始 LR 训练，未加入本项目超分、颜色校正、位移场或可见性增益模块。下表是本轮统一协议的复现结果，不是原论文表格数值。")
    paragraph("每景使用 300 帧。Cook 有 20 台训练相机，Cut 因缺少 cam04 使用 19 台；cam00 全程留出。训练为 336×252 的原始 bicubic antialias LR，最终渲染为 1344×1008。真实 HR 不参与模型训练或初始化，仅用于身份校验、预先固定可视化区域，以及训练后的评价和图像检查。4C4D 使用作者的 MASt3R 几何先验，单独披露如下。")
    heading("留出相机指标")
    rows = []
    for method in METHODS:
        row = [metadata[method]["label"] + "\n" + metadata[method]["venue"]]
        for scene in SCENES:
            values = records[method, scene]["metrics"]
            row.extend([f"{values['psnr']:.3f}", f"{values['ssim']:.4f}", f"{values['lpips']:.4f}"])
        rows.append(row)
    table(["方法及出处", "Cook\nPSNR ↑", "Cook\nSSIM ↑", "Cook\nLPIPS ↓",
           "Cut\nPSNR ↑", "Cut\nSSIM ↑", "Cut\nLPIPS ↓"], rows, [1.4, .95, .95, .95, .95, .95, .95])
    paragraph("PSNR 为逐帧全图 PSNR 的算术平均；SSIM 使用 11×11 高斯窗口、σ=1.5，去除外侧 5 像素；LPIPS 使用 AlexNet v0.1。所有主表数值均对 8 位 RGB PNG 在相同 A100 base 环境重算，未混入历史浮点渲染指标。")
    summaries = []
    for scene in SCENES:
        winners = []
        for metric, label, choose in (("psnr", "PSNR 最高", max), ("ssim", "SSIM 最高", max), ("lpips", "LPIPS 最低", min)):
            winner = choose(METHODS, key=lambda method: records[method, scene]["metrics"][metric])
            winners.append(label + "为 " + metadata[winner]["label"])
        summaries.append(LABELS[scene] + "：" + "，".join(winners))
    paragraph("；".join(summaries) + "。这是固定端点的全图均分比较，区域细节仍须结合后文图像检查。")

    document.add_page_break()
    heading("作者日程与初始化条件")
    rows = [[metadata[m]["label"], metadata[m]["native_schedule"], metadata[m]["initialization"]] for m in METHODS]
    table(["方法", "固定日程", "合法训练 LR 初始化"], rows, [1.15, 2.35, 3.15], left_columns=(0, 1, 2))
    paragraph("Wu 和 Multi4D 使用已标定训练 LR 的三角化点，初始化方式与作者原 COLMAP 输入存在差异。Yang 使用训练第 0 帧 COLMAP 稠密点；LocalDyGS 使用多个时刻的 COLMAP 稠密点；CEM 使用第 0 帧 COLMAP 稀疏点。初始化差异和作者日程差异均保留，不能据此声称严格隔离了骨干结构的因果作用。")
    paragraph("LocalDyGS/CEM 的 Cook 初始化使用 CTS 的 COLMAP 3.12dev，Cut 使用 A100 的 COLMAP 3.9.1，实际软件身份随初始化保存。历史 Yang 初始化未冻结当时的 COLMAP 二进制哈希，当前机器上的二进制身份不能追认为其历史身份。这些初始化软件差异属于本轮复现的限制。")
    paragraph("4C4D 使用作者 MAtCha posed MASt3R 预训练几何先验，只输入训练 LR。模型内部按作者设置调整工作分辨率；这不构成 SR 图像监督。原生过程允许内部平移与尺度优化，之后对齐至给定相机。最终动态重建与渲染使用 manifest 相机。")
    paragraph("4C4D 原论文面向四相机稀疏采集。本轮为统一输入，使用 Cook 的全部 20 台和 Cut 的全部 19 台合法训练相机。因此，表中 4C4D 应标为统一多视角 LR 协议适配结果；引用原论文四相机指标时须另行区分。")
    paragraph("为适配当前环境，保留相机主点与半像素约定、原生时间尺度，并修复已证实的图像类型、CUDA 头文件和进程资源兼容问题。各修复、固定上游提交和训练完整性验收均随端点保存。失败的早期部署尝试单独保留，不作为正式结果。")
    paragraph("表中日程按作者迭代编号标注。Yang 与 4C4D 的原生循环从编号 2 开始，编号 30000 末步不更新优化器，因此从头完成该端点对应 29998 次参数更新。4C4D 还保留作者不透明度衰减和点增密至编号 30000 的实际覆盖设置。本轮未调整这些原生行为。")
    paragraph("LocalDyGS 的作者入口保存点云、动态哈希场和 MLP 权重，不保存完整优化器状态。本轮已回传这些原生最终资产，可重新加载渲染；后续若继续微调，需要重新建立优化器，不能称为恢复完整训练状态。")
    heading("结论的适用范围")
    paragraph("这两个场景用于当前对照准备与开发验证，不能代表完整标准基准，也不能单凭图像均分证明三维几何或时序一致性。固定帧图像检查属于代理检查，尚未进行用户盲评或定量时序评价。")
    paragraph("后续比较本项目方法时，应使用完全相同的原始 LR、相机划分和评价定义重新运行对应方法。当前 area4 输入下的本项目结果不能直接加入本表。将 SR 教师输出作为监督的分支应另列，并披露额外先验。")

    document.add_page_break()
    heading("训练资源与可追踪记录")
    rows = []
    for method in METHODS:
        for scene in SCENES:
            gpu, minutes, memory = cost(records[method, scene]["native_endpoint_metadata"], method)
            rows.append([metadata[method]["label"], LABELS[scene], gpu, minutes, memory])
    table(["方法", "场景", "训练设备", "训练分钟", "峰值 GiB"], rows, [1.25, .6, 1.8, 1.25, 1.25], left_columns=(0, 2))
    paragraph("Wu 复用已核验的历史 PRO6000 干净端点，本轮没有重复训练。其他方法记录实际 A100 或 CTS 训练设备。显存统一换算为 GiB；各作者入口记录已分配或已保留显存，具体口径见端点索引，不能覆盖外部 Faiss 或 COLMAP。未知值如实保留。各方法计时范围不同，精确范围见端点索引，不能据本表直接宣称公平的速度或显存排序。")
    paragraph("统一评价消除了评价实现与评价设备的差异，但不消除训练硬件对优化过程的影响。细微增益需要同硬件或配对重复确认。初始化、环境构建、失败恢复和完整作业耗时在原始收据中单独保存。")
    paragraph("本地证据位于 output/dynamic_lr_baselines_20261020。comparison/comparison.json 记录十二个端点、检查点与指标身份；comparison.csv 可直接整理结果表。每个评价目录保存 per_frame.csv 和逐帧图像 SHA256。配置、补丁、环境与初始化证据随各端点回传。")
    paragraph("输入清单共 12306 个文件，两台主机均核验无缺失或哈希差异。最终汇总再次核验 3600 张 PNG 的完整性、分辨率、实际评价字节与参考 HR 身份。完整哈希保存在机器记录中，本文不手工转录。")
    heading("正式论文来源")
    for method in METHODS:
        item = metadata[method]
        label = item["label"] + "  " + item["venue"]
        para = document.add_paragraph()
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), document.part.relate_to(item["paper_url"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = label + " 正式论文"
        run.append(text)
        link.append(run)
        para._p.append(link)
        md.extend([f"[{label}]({item['paper_url']})", ""])

    for scene in SCENES:
        document.add_page_break()
        heading(LABELS[scene] + " 固定帧全图检查")
        image = a.visuals / f"{scene}_cam00_0150_whole.png"
        picture(image, 5.2)
        paragraph("cam00 帧索引 150（从 0 开始）；每幅图以半尺寸显示，完整分辨率渲染保留。", "Caption")
        md.extend([f"![{LABELS[scene]} 第150帧全图]({image.resolve()})", ""])

    document.add_page_break()
    heading("固定帧人脸细节检查")
    paragraph("图示为 cam00 的帧索引 150，索引从 0 开始。人脸区域在查看新方法结果之前由 HR 固定；128×128 区域以最近邻放大 2 倍，未锐化。人物低头，区域主要覆盖头发、额头和眼镜。另保存两景帧索引 0、150、299 的全部十二张全图与人脸对照面板。")
    for scene in SCENES:
        image = a.visuals / f"{scene}_cam00_0150_face.png"
        picture(image, 5.8)
        paragraph(LABELS[scene] + " 帧索引 150 固定区域", "Caption")
        notes = review.get("scene_notes", {}).get(scene)
        if notes:
            paragraph(notes)
        md.extend([f"![{LABELS[scene]} 第150帧]({image.resolve()})", ""])

    a.out.parent.mkdir(parents=True, exist_ok=True)
    for element in (document.element, document.styles.element):
        for border in list(element.xpath(".//w:pBdr")):
            border.getparent().remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.line_spacing = Pt(10)
    footer.add_run("第 ").font.size = Pt(8)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    footer.add_run(" 页").font.size = Pt(8)
    docx = a.out.with_suffix(".docx")
    assert not docx.exists(), "Preserve an existing report; use an explicit new output stem"
    document.save(docx)
    a.out.with_suffix(".md").write_text("\n".join(md) + "\n")
    receipt = dict(status="authored_pending_DOCX_render_inspection", docx=dict(path=str(docx.resolve()), sha256=sha(docx)),
                   comparison_sha256=sha(a.comparison), metadata_sha256=sha(a.metadata), visual_review_sha256=sha(a.review),
                   endpoint_count=12, predicted_PNGs_verified=3600)
    a.out.with_suffix(".source.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(json.dumps(receipt))


if __name__ == "__main__":
    main()
