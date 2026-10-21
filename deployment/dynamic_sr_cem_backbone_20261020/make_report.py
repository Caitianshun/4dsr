"""Create the CEM area4 Word/Markdown report after complete scientific inputs.

Use the Codex bundled Python runtime for authoring. The caller must run the
artifact-operation marker immediately before authoring, then render the DOCX
and inspect every page separately. This helper never runs either operation.
--check-only is stdlib-only, creates no output, and never imports python-docx.
No model, training log, CUDA or network operation is used.

Actual root review contract:
 schema: CEM_area4_actual_visual_review_v1
 status: reviewed_by_agent; reviewer: nonempty; human_acceptance: false
 comparison / diagnostic_receipt: {path: absolute, sha256: ...}
 conclusions: [actual scientific interpretation strings]
 scene_notes: {cook_spinach: {cam00:[...],cam01:[...]}, cut_roasted_beef: {...}}
 inspected_panels: [{path:absolute,sha256:...,reviewer:...,observations:[...]}]
The review must cover the eight inserted frame150 raw whole/face panels and
any available two inserted Wu/CEM matched-arm face panels. Additional actual
reviews of frames0/299 are permitted and preserved, not invented or auto-rated.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import struct
import sys

ROOT = Path(__file__).resolve().parents[2]
BATCH = ROOT / "output/dynamic_sr_cem_backbone_20261020"
SCENES = ("cook_spinach", "cut_roasted_beef")
LABELS = dict(cook_spinach="Cook", cut_roasted_beef="Cut")
ARMS = ("areaL0", "naive", "areaT", "areaTm", "areaTm_w4", "areaOracle")
CEM_ARMS = ("native_parent", *ARMS, "HRref")
ARM_LABELS = dict(native_parent="原生 LR 父模型", areaL0="areaL0", naive="naive", areaT="areaT", areaTm="areaTm",
                  areaTm_w4="areaTm + w4", areaOracle="areaOracle", HRref="原生 HRref")
WU_KEYS = dict(zip(ARMS, ("areaL0", "naive", "areaT", "areaTm", "w4", "oracle")))
METRICS = ("psnr", "ssim", "lpips")
BUNDLED_PYTHON = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3"
ENV_KEYS = ("torch", "numpy", "opencv", "gpu", "evaluator_sha256", "helper_source_sha256", "lpips_state_dict_sha256")


def require(value, message):
    if not value:
        raise ValueError(message)


def read(path):
    return json.loads(Path(path).read_text())


def sha(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(8 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def identity(path):
    path = Path(path).resolve()
    return dict(path=str(path), sha256=sha(path), bytes=path.stat().st_size)


def check_identity(item, expected_path=None):
    path = Path(item["path"])
    require(path.is_absolute(), "Report/review identities must use absolute local paths")
    path = path.resolve()
    if expected_path is not None:
        require(path == Path(expected_path).resolve(), "Identity path differs: " + str(path))
    before = path.stat()
    value = identity(path)
    after = path.stat()
    require((before.st_ino, before.st_size, before.st_mtime_ns) == (after.st_ino, after.st_size, after.st_mtime_ns), "Source changed during verification")
    require(value["sha256"] == item["sha256"], "Source SHA differs: " + str(path))
    if "bytes" in item:
        require(value["bytes"] == item["bytes"], "Source bytes differ")
    return value


def texts(values, label):
    require(isinstance(values, list) and values and all(isinstance(x, str) and x.strip() for x in values), label + " must contain actual nonempty text entries")
    return values


def png_size(path):
    with Path(path).open("rb") as stream:
        header = stream.read(29)
    require(len(header) == 29 and header[:8] == b"\x89PNG\r\n\x1a\n" and header[12:16] == b"IHDR", "Not a PNG panel")
    width, height = struct.unpack(">II", header[16:24])
    require(width > 0 and height > 0, "Invalid PNG dimensions")
    return width, height


def output_gate(stem):
    stem = Path(stem).resolve()
    require(stem.is_relative_to((ROOT / "docs").resolve()), "Output stem must be inside this project's docs directory")
    require(stem.suffix not in {".docx", ".md", ".json"}, "--out takes a stem without a document extension")
    for suffix in (".docx", ".md", ".source.json"):
        require(not Path(str(stem) + suffix).exists(), "Preserve previous report: " + str(stem) + suffix)
    require(not Path(str(stem) + ".report_assets").exists(), "Preserve previous report display assets")
    return stem


def serial_io_amendment(args, rows):
    """Bind the real engineering amendment; never parse a training log."""
    amendment_path = getattr(args, "protocol_amendment", BATCH / "protocol_amendment_serial_HR_IO_v2.json").resolve()
    qa_path = getattr(args, "io_cpu_qa", BATCH / "admission/serial_io_CPU_QA_CTS_v1.json").resolve()
    amendment, qa = read(amendment_path), read(qa_path)
    require(amendment["schema"] == "CEM_area4_predeclared_matrix_engineering_amendment_v1" and amendment["scope"] == "independent native training-HR references only", "Wrong HR IO protocol amendment")
    require(amendment["failed_runs_preserved"] is True and amendment["selection_on_validation_metrics"] is False, "HR recovery archive/selection boundary absent")
    require("n_job=1" in amendment["change"] and "preserve original native source bytes" in amendment["change"], "Wrong IO adaptation")
    require(set(amendment["new_output_runs"]) == {f"runs/{scene}_native_HRref_v3" for scene in SCENES}, "Unexpected amended HR endpoints")
    evidence = {}
    for item in amendment["evidence"]:
        path = Path(item["path"])
        require(not path.is_absolute() and (ROOT / path).resolve().is_relative_to(BATCH), "Amendment evidence outside batch")
        require(str(path) not in evidence, "Duplicate amendment evidence")
        evidence[str(path)] = check_identity(dict(path=str((ROOT / path).resolve()), sha256=item["sha256"]))
    host_path = "output/dynamic_sr_cem_backbone_20261020/admission/HRref_SIGKILL_host_evidence_v1.json"
    failed_path = "output/dynamic_sr_cem_backbone_20261020/runs/cook_spinach_native_HRref_v2/status.json"
    require(host_path in evidence and failed_path in evidence, "Real kernel/failed execution evidence missing")
    host, failed = read(evidence[host_path]["path"]), read(evidence[failed_path]["path"])
    kernel = host["cts"]["kernel_window"]["stdout"]
    require("global_oom" in kernel and "Killed process 445330 (python)" in kernel and "anon-rss:63185968kB" in kernel, "Host OOM is not confirmed by archived kernel evidence")
    require(failed["status"] == "failed" and failed["scene"] == "cook_spinach" and failed["arm"] == "native_HRref", "Wrong archived failure")
    failed_stages = [x for x in failed["stages"] if x["name"] == "native_stage1"]
    require(len(failed_stages) == 1 and failed_stages[0]["returncode"] == -9, "Wrong archived stage failure")
    seconds = failed_stages[0]["seconds"]
    require(type(seconds) in (int, float) and math.isfinite(seconds) and seconds > 0, "Missing failed-stage cost")
    failed_run = Path(evidence[failed_path]["path"]).parent
    require(not (failed_run / "stage1.json").exists() and not any((failed_run / "model" / f"chkpnt{x}.pth").exists() for x in (30000, 40000)), "Failure archive now contains a different completed/resumable endpoint")
    wrapper = ROOT / "experiments/dynamic_sr_cem_area_20261020/native_stage_serial_io.py"
    native_entry = ROOT / "experiments/dynamic_sr_cem_area_20261020/native_stage.py"
    require(qa["schema"] == "CEM_native_HR_serial_io_CPU_QA_v1" and qa["status"] == "passed" and qa["CUDA_initialized"] is False, "Actual CTS CPU IO QA required")
    require(qa["joblib_version"] == "1.6.0" and qa["wrapper"]["sha256"] == sha(wrapper) and qa["native_entry_sha256"] == sha(native_entry), "CTS IO QA/source identity differs")
    require(qa["joblib_parallel_source"]["sha256"] == sha(BATCH / "admission/CTS_joblib_1_6_source_v1/parallel.py"), "CTS joblib source identity differs")
    require(qa["same_seed_camera_order_equal"] is True and qa["main_process_Python_NumPy_Torch_CPU_RNG_unchanged_by_decoding"] is True, "IO camera order/RNG QA failed")
    require(qa["generator_before_next"] == dict(task_iterator_advanced=0, decoded=0) and qa["generator_after_one_next_and_wait"] == dict(task_iterator_advanced=1, decoded=1), "IO generator is not demand driven")
    require(len(qa["image_evidence"]) == 6 and all(x["bitwise_equal"] is True for x in qa["image_evidence"]), "Real-image bitwise QA missing")
    require(qa["native_model_imported"] is False and qa["native_training_executed"] is False, "CPU QA scope differs")
    io_sources = []
    for scene in SCENES:
        row = rows["CEM", scene, "HRref"]
        require(Path(row["run"]).name == f"{scene}_native_HRref_v3", "Final HRref is not the amended fresh run")
        stages = row["original_training_receipts"]
        require(len(stages) == 2, "Both amended HR stages required")
        for number, stage in enumerate(stages, 1):
            adaptation = stage["io_adaptation"]
            require(stage["stage"] == number and stage["entry_sha256"] == sha(native_entry), "Original native entry/stage identity changed")
            require(adaptation["schema"] == "CEM_native_HR_serial_io_v1" and adaptation["status"] == "completed" and adaptation["workers"] == adaptation["n_job"] == 1 and adaptation["no_core_modified"] is True, "Final HR endpoint lacks completed serial IO evidence")
            require(adaptation["wrapper"]["sha256"] == qa["wrapper"]["sha256"] and adaptation["native_entry"]["sha256"] == qa["native_entry_sha256"], "Final HR IO sources differ from CPU QA")
            io_path = Path(row["run"]) / f"stage{number}.io.json"
            io_id = check_identity(dict(path=str(io_path.resolve()), sha256=adaptation["receipt"]["sha256"]))
            trace = read(io_path)
            require(trace["status"] == "completed" and trace["schema"] == adaptation["schema"] and trace["no_core_modified"] is True and trace["workers"] == trace["n_job"] == 1, "Independent IO trace incomplete")
            require(trace["wrapper"] == adaptation["wrapper"] and trace["native_entry"] == adaptation["native_entry"], "Independent IO trace source differs")
            require(trace["frozen_original_method"] == adaptation["original_method"], "Independent IO original method differs")
            require(all(adaptation["original_method"][key] == qa["frozen_original_method"][key] for key in ("scene_sha256", "method_ast_sha256", "reader_ast_sha256", "protocol_receipt_sha256")), "Final HR original method differs from CTS CPU QA")
            calls, count = trace["calls"], adaptation["completed_calls"]
            require(type(count) is int and count >= 1 and len(calls) == 2 * count, "Incomplete HR IO call trace")
            for before, after in zip(calls[::2], calls[1::2]):
                require(before["phase"] == "before_original_call" and after["phase"] == "after_original_call" and before["n_job"] == before["forced_n_job"] == 1, "HR IO call did not use worker one")
                require({k: v for k, v in before.items() if k != "phase"} == {k: v for k, v in after.items() if k not in ("phase", "camera_count")}, "Unpaired HR IO calls")
            io_sources.append(io_id)
    return dict(amendment=amendment, failed_stage_seconds=seconds, failed_job_seconds=failed["elapsed_seconds"],
        sources=dict(protocol_amendment=identity(amendment_path), actual_CTS_serial_IO_CPU_QA=identity(qa_path),
            serial_IO_wrapper=identity(wrapper), original_native_entry=identity(native_entry)),
        evidence_sources=list(evidence.values()), final_HR_IO_receipts=io_sources)



def report_tile_pages(item):
    """Read-only display contract, preserving scientific tile order and sources."""
    panel, camera, role = item["panel"], item["camera"], item["role"]
    require(item["region"] in {"whole", "face"} and camera in {"cam00", "cam01"}, "Wrong report image scope")
    tiles = panel["tiles"]
    keys = [tile["key"] for tile in tiles]
    if role == "main_raw":
        expected = ["GT", "LR_bicubic"] + (["teacher_original", "teacher_matched"] if camera == "cam01" else []) + list(CEM_ARMS)
    else:
        require(role == "Wu_CEM_paired_raw" and camera == "cam00" and item["region"] == "face", "Unknown report layout role")
        expected = ["GT", *("Wu_" + arm for arm in ARMS), "GT", *ARMS]
    require(keys == expected, "Report tile order/key differs from the fixed diagnostic protocol")
    roi = panel["fixed_face_xyxy"]
    require(len(roi) == 4 and all(type(x) is int for x in roi) and 0 <= roi[0] < roi[2] <= 1344 and 0 <= roi[1] < roi[3] <= 1008, "Invalid fixed report ROI")
    for tile in tiles:
        require(not tile.get("missing"), "Missing source cannot enter a report grid")
        key = tile["key"]
        if key == "GT":
            expected_role = "clean_HR_reference"
        elif key == "LR_bicubic":
            expected_role = "input_visual_reference_not_a_reconstruction_method"
        elif key in {"teacher_original", "teacher_matched"}:
            require(camera == "cam01", "Teacher cannot appear in the held-out camera")
            expected_role = "training_camera_teacher_reference_only"
            check_identity(tile["teacher_index"])
        elif key.startswith("Wu_"):
            expected_role = "historically_selected_Wu_development_reference_no_posthoc_cull"
        else:
            expected_role = "heldout_raw_reconstruction" if camera == "cam00" else "training_camera_raw_reconstruction"
        require(tile["role"] == expected_role, "Corrected/unknown source role cannot enter a raw report grid")
        operation = "LR_bicubic_visual_reference" if key == "LR_bicubic" else "none"
        require(tile["source_operation"] == operation, "Unknown report source image operation")
        require(tile["display_crop_xyxy"] == (roi if item["region"] == "face" else None), "Report crop differs from the fixed diagnostic ROI")
        check_identity(tile["source"])
        require(png_size(tile["source"]["path"]) == ((336, 252) if key == "LR_bicubic" else (1344, 1008)), "Wrong report source PNG dimensions")
    if role == "Wu_CEM_paired_raw":
        require(tiles[0]["source"] == tiles[7]["source"], "Wu/CEM GT references differ")
        # Same three method columns, top Wu and bottom CEM, in author arm order.
        return [dict(number=number + 1, indices=list(range(1 + start, 4 + start)) + list(range(8 + start, 11 + start)),
                     method_columns=list(ARMS[start:start + 3]), GT_reference=tiles[0]["source"])
                for number, start in enumerate((0, 3))]
    limit = 6 if item["region"] == "face" else 12
    return [dict(number=number + 1, indices=list(range(start, min(start + limit, len(tiles)))))
            for number, start in enumerate(range(0, len(tiles), limit))]


def report_tile_label(item, tile):
    key = tile["key"]
    labels = {"GT": "GT HR", "LR_bicubic": "LR bicubic · 输入参照", "teacher_original": "Teacher original · 训练伪HR",
              "teacher_matched": "Teacher matched · 训练伪HR", "native_parent": "CEM LR parent · raw",
              "HRref": "CEM HRref · 训练HR诊断"}
    if key in labels:
        return labels[key]
    if key.startswith("Wu_"):
        return "Wu " + key[3:].replace("areaTm_w4", "areaTm + w4") + " · 开发端点 raw"
    return "CEM " + key.replace("areaTm_w4", "areaTm + w4") + (" · 训练HR诊断 raw" if key == "areaOracle" else " · raw")


def report_crop(item, index, directory):
    """Pillow-only display pixels. Called after complete preflight by author()."""
    from PIL import Image, ImageDraw
    tile = item["panel"]["tiles"][index]
    source = check_identity(tile["source"])
    with Image.open(source["path"]) as original:
        require(original.format == "PNG" and original.mode == "RGB", "Report source must be an RGB PNG")
        expected = (336, 252) if tile["key"] == "LR_bicubic" else (1344, 1008)
        require(original.size == expected, "Report source decoded dimensions differ")
        image = original.copy()
    check_identity(tile["source"])
    operation = tile["source_operation"]
    if operation == "LR_bicubic_visual_reference":
        # Identical PIL operation to frozen diagnostics Assets.image().
        image = image.resize((1344, 1008), Image.Resampling.BICUBIC)
    roi = item["panel"]["fixed_face_xyxy"]
    if item["region"] == "face":
        image = image.crop(tuple(roi))
        processing = "fixed ROI crop; no enhancement or display resampling; Word fits original crop aspect"
    else:
        ImageDraw.Draw(image).rectangle(tuple(roi), outline="#ffee00", width=3)
        processing = "whole original image plus the frozen diagnostic yellow ROI outline; no enhancement"
    name = f"{item['scene']}_{item['camera']}_0150_{item['region']}_{item['role']}_t{index:02d}_{tile['key']}.png"
    path = directory / name
    require(not path.exists(), "Preserve previous report crop")
    with path.open("xb") as stream:
        image.save(stream, format="PNG")
    return dict(source=source, image=identity(path), key=tile["key"], role=tile["role"],
                original_panel=dict(path=item["panel"]["path"], sha256=item["panel"]["sha256"]),
                original_tile_index=index, source_operation=operation,
                fixed_crop_xyxy=list(roi) if item["region"] == "face" else None,
                processing=processing, width=image.width, height=image.height,
                decoded_RGB_sha256=hashlib.sha256(image.tobytes()).hexdigest())


def preflight(args):
    """All preconditions are read-only and precede importing docx or mkdir."""
    stem = output_gate(args.out)
    comparison_path = args.comparison.resolve()
    comparison = read(comparison_path)
    require(comparison["schema"] == "CEM_Wu_area4_final_receipt_comparison_v1", "Wrong comparison schema")
    require(comparison["status"] == "complete" and comparison["complete"] is True and comparison["verified_endpoints"] == comparison["expected_endpoints"] == 28, "A formal report requires complete28 results")
    require(not comparison["missing_or_invalid"] and set(comparison["selected_scenes"]) == set(SCENES), "Partial scene/results cannot be reported as complete")
    expected = {("CEM", scene, arm) for scene in SCENES for arm in CEM_ARMS} | {("Wu", scene, arm) for scene in SCENES for arm in ARMS}
    rows = {(row["system"], row["scene"], row["arm"]): row for row in comparison["endpoints"]}
    require(set(rows) == expected and len(comparison["endpoints"]) == 28, "Missing/duplicated/wrong endpoint keys")
    require(comparison["acceptance"]["input_result_matrix_complete"] is True and comparison["acceptance"]["human_visual_acceptance"] is False and comparison["acceptance"]["pure_backbone_causal_attribution"] is False, "Unsupported acceptance/attribution in comparison")
    complete_gate = read(comparison_path.parent / "complete_matrix_gate.json")
    require(complete_gate["schema"] == "CEM_Wu_area4_complete_matrix_gate_v1" and complete_gate["status"] == "complete" and complete_gate["endpoints"] == 28, "Missing formal summarizer completion gate")
    check_identity(complete_gate["comparison"], comparison_path)
    for row in rows.values():
        require(row["status"] == "verified" and row["heldout_training"] is False, "Unverified/leaking endpoint")
        metric = row["metrics"]
        require(all(type(metric["values"][key]) in (int, float) and math.isfinite(metric["values"][key]) for key in METRICS), "Nonfinite/missing main metric")
        check_identity(metric["summary"])
        check_identity(metric["per_frame"])
        original = read(metric["summary"]["path"])
        require(original == metric["original_summary"] and original["test"] == metric["values"], "Embedded/source main metrics differ")
        require(original["status"] == "completed" and original["schema"] == "dynamic_lr_shared_HR_image_metrics_v1" and original["frames"] == 300 and original["camera"] == "cam00", "Wrong public metric protocol")
        require(original["device"] == "cuda" and "A100" in original["gpu"] and {key: original[key] for key in ENV_KEYS} == comparison["common_metric_environment"], "Metrics not from same A100 environment")
        require(original["input_representation"] == "8-bit rendered RGB PNG for every method", "Raw float/8bit comparison mixing forbidden")
        require(original["manifest_sha256"] == row["manifest"]["sha256"], "Metric/endpoint manifest differs")
    diagonal = {(x["system"], x["scene"], x["contrast"]): x for x in comparison["paired_deltas"]}
    for system, contrasts in [("CEM", ["areaT-naive", "areaTm-areaT", "w4-areaTm", "Oracle-areaTm", "nativeHR-parent"]), ("Wu", ["areaT-naive", "areaTm-areaT", "w4-areaTm", "Oracle-areaTm"])]:
        for scene in SCENES:
            for contrast in contrasts:
                item = diagonal[system, scene, contrast]
                require(item["status"] == "paired_available", "Missing paired delta")
                a, b = rows[system, scene, item["target"]], rows[system, scene, item["source"]]
                require(item["delta_target_minus_source"] == {key: a["metrics"]["values"][key] - b["metrics"]["values"][key] for key in METRICS}, "Paired delta differs")

    receipt_path = args.diagnostics.resolve() / "visual_diagnostic_receipt.json"
    visual = read(receipt_path)
    require(visual["status"] == "complete_fixed_image_diagnostics" and visual["final_gate"] is True and not visual["missing_required_endpoints"], "Formal fixed-image diagnostics required")
    require(all(visual[key] is False for key in ("human_acceptance", "temporal_acceptance", "scientific_acceptance", "blinded")), "Image input gate is not human/scientific acceptance")
    fixed = check_identity(visual["fixed_protocol"])
    protocol = read(fixed["path"])
    require(protocol["status"] == "fixed_before_new_CEM_results" and protocol["frames"] == [0, 150, 299] and protocol["cameras"] == dict(cam00="heldout", cam01="training_diagnostic"), "ROI/frame/camera protocol changed")
    image_gate = read(args.diagnostics / "complete_image_inputs_gate.json")
    require(image_gate["status"] == "all_required_returned_endpoints_and_fixed_image_inputs_verified" and image_gate["human_acceptance"] is False, "Diagnostic completeness gate absent")
    check_identity(image_gate["diagnostic_receipt"], receipt_path)
    diagnostic_endpoints = {}
    for endpoint in visual["endpoint_receipts"]:
        require(not endpoint.get("missing"), "Missing visual endpoint")
        system = "Wu" if endpoint["kind"] == "Wu_development_reference" else "CEM"
        key = endpoint["key"][3:] if system == "Wu" else endpoint["key"]
        lookup_key = (system, endpoint["scene"], key)
        require(lookup_key not in diagnostic_endpoints, "Duplicate visual endpoint")
        require(lookup_key in rows, "Unknown visual endpoint")
        row = rows[lookup_key]
        require(endpoint["checkpoint_sha256"] == row["checkpoint"]["sha256"] and Path(endpoint["run"]).resolve() == Path(row["run"]).resolve(), "Visual/main-table endpoint differs")
        for item in endpoint["receipt_identities"]:
            check_identity(item)
        diagnostic_endpoints[lookup_key] = endpoint
    require(set(diagnostic_endpoints) == expected, "Formal visuals must include the same28 endpoints")
    all_panels = {}
    for panel in visual["panels"]:
        path = Path(panel["path"])
        require(path.is_absolute() and path.resolve().is_relative_to(args.diagnostics.resolve()), "Panel outside declared diagnostic directory")
        key = str(path.resolve())
        require(key not in all_panels, "Duplicate panel path")
        all_panels[key] = panel
    image_sources = {item["path"]: item for item in visual["image_sources"]}
    selected = []
    for scene in SCENES:
        for camera in ("cam00", "cam01"):
            for region in ("whole", "face"):
                path = args.diagnostics.resolve() / f"{scene}_{camera}_0150_{region}_raw.png"
                require(str(path) in all_panels, "Missing frame150 raw panel: " + path.name)
                panel = all_panels[str(path)]
                require(panel["region"] == region and panel["fixed_face_xyxy"] == protocol["face_xyxy"][scene][camera], "Panel ROI changed")
                keys = {tile["key"] for tile in panel["tiles"]}
                expected_keys = {"GT", "LR_bicubic", *CEM_ARMS}
                if camera == "cam01":
                    expected_keys |= {"teacher_original", "teacher_matched"}
                require(keys == expected_keys and len(keys) == len(panel["tiles"]), "Missing/extra main-panel method or illegal cam00 teacher")
                selected.append(dict(scene=scene, camera=camera, region=region, panel=panel, role="main_raw"))
    paired = []
    for scene in SCENES:
        path = args.diagnostics.resolve() / f"{scene}_cam00_0150_face_Wu_CEM_matched_arms.png"
        if str(path) in all_panels:
            panel = all_panels[str(path)]
            require(panel["region"] == "face" and panel["fixed_face_xyxy"] == protocol["face_xyxy"][scene]["cam00"], "Wu paired ROI changed")
            require([tile["key"] for tile in panel["tiles"]] == ["GT", *("Wu_" + arm for arm in ARMS), "GT", *ARMS], "Wu/CEM arms not aligned")
            paired.append(dict(scene=scene, camera="cam00", region="face", panel=panel, role="Wu_CEM_paired_raw"))
    require(len(paired) in (0, 2), "Only one scene has a paired Wu/CEM panel")
    selected += paired
    for item in selected:
        panel = item["panel"]
        check_identity(panel)
        require(png_size(panel["path"]) == (panel["width"], panel["height"]), "Panel PNG dimensions differ")
        require("PARTIAL" not in panel.get("title", ""), "Partial panel cannot enter report")
        for tile in panel["tiles"]:
            require(not tile.get("missing"), "Missing tile cannot enter report")
            source = tile["source"]
            require(source["path"] in image_sources and source["sha256"] == image_sources[source["path"]]["sha256"], "Tile/source receipt differs")
            check_identity(source)
            if tile["key"] not in {"GT", "LR_bicubic", "teacher_original", "teacher_matched"}:
                require(tile["role"] in {"heldout_raw_reconstruction", "training_camera_raw_reconstruction", "historically_selected_Wu_development_reference_no_posthoc_cull"}, "Corrected image cannot replace raw reconstruction")

    for item in selected:
        report_tile_pages(item)

    review_path = args.review.resolve()
    review = read(review_path)
    require(review["schema"] == "CEM_area4_actual_visual_review_v1" and review["status"] == "reviewed_by_agent" and review["human_acceptance"] is False, "Actual agent visual review required; do not claim human acceptance")
    require(isinstance(review["reviewer"], str) and review["reviewer"].strip(), "Reviewer missing")
    check_identity(review["comparison"], comparison_path)
    check_identity(review["diagnostic_receipt"], receipt_path)
    texts(review["conclusions"], "Root conclusions")
    require(set(review["scene_notes"]) == set(SCENES), "Actual observations required for both scenes")
    for scene in SCENES:
        for camera in ("cam00", "cam01"):
            texts(review["scene_notes"][scene][camera], scene + " " + camera + " observations")
    inspected = {}
    for item in review["inspected_panels"]:
        path = Path(item["path"])
        require(path.is_absolute() and str(path.resolve()) in all_panels, "Reviewed image not present in formal diagnostics")
        name = str(path.resolve())
        require(name not in inspected, "Duplicate inspected panel")
        require(item["sha256"] == all_panels[name]["sha256"], "Review/diagnostic panel SHA differs")
        require(isinstance(item["reviewer"], str) and item["reviewer"].strip(), "Panel reviewer missing")
        texts(item["observations"], "Actual panel observations")
        check_identity(item)
        inspected[name] = item
    require(all(item["panel"]["path"] in inspected for item in selected), "An inserted panel has not been actually reviewed")
    io = serial_io_amendment(args, rows)
    return dict(stem=stem, comparison=comparison, rows=rows, deltas=diagonal, visual=visual, protocol=protocol, io_amendment=io,
        review=review, selected=selected, inspected=inspected, sources=dict(comparison=identity(comparison_path),
        comparison_complete_gate=identity(comparison_path.parent / "complete_matrix_gate.json"),
        diagnostic_receipt=identity(receipt_path), diagnostic_complete_gate=identity(args.diagnostics / "complete_image_inputs_gate.json"),
        fixed_protocol=fixed, actual_review=identity(review_path), generator=identity(__file__), **io["sources"]))


def metric_rows(state, system):
    result = []
    for arm in CEM_ARMS if system == "CEM" else ARMS:
        row = [ARM_LABELS[arm]]
        for scene in SCENES:
            values = state["rows"][system, scene, arm]["metrics"]["values"]
            row += [f"{values['psnr']:.3f}", f"{values['ssim']:.4f}", f"{values['lpips']:.4f}"]
        result.append(row)
    return result


def delta_rows(state, system):
    result = []
    contrasts = ["areaT-naive", "areaTm-areaT", "w4-areaTm", "Oracle-areaTm"] + (["nativeHR-parent"] if system == "CEM" else [])
    for contrast in contrasts:
        row = [contrast]
        for scene in SCENES:
            values = state["deltas"][system, scene, contrast]["delta_target_minus_source"]
            row += [f"{values['psnr']:+.3f}", f"{values['ssim']:+.4f}", f"{values['lpips']:+.4f}"]
        result.append(row)
    return result


def resource_values(endpoint):
    resource = endpoint["resources"]
    gpu = resource.get("training_GPU")
    if isinstance(gpu, list):
        gpu = " / ".join(dict.fromkeys(str(x) for x in gpu if x)) or "未知"
    if gpu:
        gpu = str(gpu).replace("NVIDIA GeForce RTX 3090", "RTX 3090").replace("NVIDIA A100-SXM4-40GB", "A100 40GB").replace("NVIDIA RTX PRO 6000 Blackwell Workstation Edition", "PRO6000")
    else:
        gpu = "未知"
    seconds = resource.get("training_seconds")
    seconds_text = "未知" if seconds is None else f"{seconds:.2f}"
    if endpoint["system"] == "Wu":
        peak, kind = resource.get("peak_allocated_bytes"), "allocated"
    else:
        peak, kind = resource.get("torch_peak_reserved_bytes"), "reserved"
    peak_text = "未知" if peak is None else f"{peak / 2**30:.3f}"
    budget = "SR 6000" if endpoint["arm"] in ARMS else "40k→总50k"
    return [LABELS[endpoint["scene"]] + " / " + ARM_LABELS[endpoint["arm"]], gpu, seconds_text, peak_text, kind, budget]


def author(state):
    """Only called after preflight; does not run marker or render acceptance."""
    require(Path(sys.executable).resolve() == BUNDLED_PYTHON.resolve(), "Author with the Codex bundled Python runtime resolved by load_workspace_dependencies")
    import docx
    require(Path(docx.__file__).resolve().is_relative_to(BUNDLED_PYTHON.parent.parent.resolve()), "Use bundled python-docx packages")
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    width = 7.1
    for style in document.styles:
        if hasattr(style, "font"):
            style.font.name = "Noto Sans CJK SC"
            style.font.color.rgb = RGBColor(0, 0, 0)
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            for name in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + name), "Noto Sans CJK SC")
                fonts.attrib.pop(qn("w:" + name + "Theme"), None)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = Pt(16.5)
    normal.paragraph_format.space_after = Pt(6)
    for name, size, line in [("Title", 21, 28), ("Heading 1", 14, 21), ("Heading 2", 12, 18), ("Caption", 9, 13)]:
        style = document.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.line_spacing = Pt(line)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.space_before = Pt(8) if name.startswith("Heading") else Pt(0)
        style.paragraph_format.keep_with_next = name.startswith("Heading")
    document.styles["Title"].font.underline = False
    md = []

    def paragraph(text, style=None):
        require(isinstance(text, str) and text.strip(), "Never write empty scientific placeholder")
        para = document.add_paragraph(text, style)
        md.extend([text, ""])
        return para

    def heading(text, page=False):
        require(not any(x in text for x in ":：—/"), "Use plain descriptive section titles")
        if page:
            document.add_page_break()
        document.add_heading(text, 1)
        md.extend(["## " + text, ""])

    def table(headers, rows, widths, left=(0,)):
        require(len(headers) == len(widths) and abs(sum(widths) - width) < .01, "Deliberate table widths must fit page")
        grid = document.add_table(rows=1, cols=len(headers))
        grid.autofit = False
        for column, col_width in zip(grid.columns, widths):
            column.width = Inches(col_width)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            rule = OxmlElement("w:" + edge)
            for name, value in dict(val="single", sz="4", color="D9D9D9").items():
                rule.set(qn("w:" + name), value)
            borders.append(rule)
        grid._tbl.tblPr.append(borders)
        for index, values in enumerate([headers, *rows]):
            require(len(values) == len(headers), "Wrong table column count")
            row = grid.rows[0] if index == 0 else grid.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for i, (cell, value, col_width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Inches(col_width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for edge, twips in dict(top=55, bottom=55, left=70, right=70).items():
                    margin = OxmlElement("w:" + edge)
                    margin.set(qn("w:w"), str(twips));margin.set(qn("w:type"), "dxa");margins.append(margin)
                props.append(margins)
                if index == 0:
                    shading = OxmlElement("w:shd");shading.set(qn("w:fill"), "E7EFF7");props.append(shading)
                cell.text = str(value)
                for para in cell.paragraphs:
                    para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(1)
                    para.paragraph_format.line_spacing = Pt(14)
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT if i in left else WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.bold = index == 0
                        run.font.color.rgb = RGBColor(0, 0, 0)
        repeated = OxmlElement("w:tblHeader");grid.rows[0]._tr.get_or_add_trPr().append(repeated)
        document.add_paragraph().paragraph_format.space_after = Pt(2)
        md.extend(["| " + " | ".join(str(v).replace("\n", "<br>") for v in headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"])
        for values in rows:
            md.append("| " + " | ".join(str(v).replace("\n", "<br>") for v in values) + " |")
        md.append("")

    display_directory = Path(str(state["stem"]) + ".report_assets")
    display_directory.mkdir(parents=True, exist_ok=False)
    state["report_image_ledger"] = []

    def report_grid(item, page):
        panel = item["panel"]
        grid = document.add_table(rows=0, cols=3)
        grid.autofit = False
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            rule = OxmlElement("w:" + edge)
            for name, value in dict(val="single", sz="4", color="D9D9D9").items():
                rule.set(qn("w:" + name), value)
            borders.append(rule)
        grid._tbl.tblPr.append(borders)
        for column in grid.columns:
            column.width = Inches(width / 3)
        for offset in range(0, len(page["indices"]), 3):
            row = grid.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for slot, index in enumerate(page["indices"][offset:offset + 3]):
                cell = row.cells[slot]
                cell.width = Inches(width / 3)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                margins = OxmlElement("w:tcMar")
                for edge, twips in dict(top=0, bottom=45, left=35, right=35).items():
                    node = OxmlElement("w:" + edge)
                    node.set(qn("w:w"), str(twips));node.set(qn("w:type"), "dxa");margins.append(node)
                cell._tc.get_or_add_tcPr().append(margins)
                tile = panel["tiles"][index]
                derived = report_crop(item, index, display_directory)
                max_width, max_height = 2.18, (2.18 if item["region"] == "face" else 1.635)
                scale = min(max_width / derived["width"], max_height / derived["height"])
                shown_width, shown_height = derived["width"] * scale, derived["height"] * scale
                para = cell.paragraphs[0]
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.line_spacing = 1
                para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(0)
                run = para.add_run()
                run.add_picture(derived["image"]["path"], width=Inches(shown_width), height=Inches(shown_height))
                for node in run._r.xpath(".//wp:docPr"):
                    node.set("descr", report_tile_label(item, tile) + "; fixed ROI; original PNG SHA " + derived["source"]["sha256"])
                label = cell.add_paragraph()
                label.alignment = WD_ALIGN_PARAGRAPH.CENTER
                label.paragraph_format.space_before = Pt(2)
                label.paragraph_format.space_after = Pt(0)
                label.paragraph_format.line_spacing = Pt(12)
                label.add_run(report_tile_label(item, tile)).font.size = Pt(10)
                derived.update(scene=item["scene"], camera=item["camera"], frame=150, region=item["region"],
                               page_group=page["number"], column=slot + 1,
                               label=report_tile_label(item, tile), label_points=10,
                               displayed_width_inches=shown_width, displayed_height_inches=shown_height)
                state["report_image_ledger"].append(derived)
        scope = "留出新视角 raw" if item["camera"] == "cam00" else "训练视角 raw 拟合诊断"
        paragraph(f"{LABELS[item['scene']]} {item['camera']} 帧150 {scope}。图下为Word原生10pt标签；固定来源原图只进行显示裁剪和等比例排版。完整科学原板与逐图来源SHA仍保留。", "Caption")
        md.extend([f"科学原板 [{Path(panel['path']).name}]({panel['path']})", ""])
        md.extend(["| 方法 | 报告显示图 |", "| --- | --- |"])
        for derived in state["report_image_ledger"][-len(page["indices"]):]:
            md.append("| " + derived["label"] + " | [固定来源图](" + derived["image"]["path"] + ") |")
        md.append("")

    title = "CEM骨干动态超分固定端点对比"
    document.add_paragraph(title, "Title");md.extend(["# " + title, ""])
    paragraph("本报告比较Cook与Cut两景的CEM原生LR父模型、六个超分对照臂、独立原生HR参考模型，以及Wu骨干对应的六个历史固定端点。所有主指标均在统一A100环境对300帧留出相机的8位RGB PNG重新评价。报告回答当前证据是否支持调整骨干；图像观察与科学判断采用实际审阅记录，未由指标自动推断。")
    for text in state["review"]["conclusions"]:
        paragraph(text)
    heading("CEM留出相机指标")
    headers = ["端点", "Cook\nPSNR ↑", "Cook\nSSIM ↑", "Cook\nLPIPS ↓", "Cut\nPSNR ↑", "Cut\nSSIM ↑", "Cut\nLPIPS ↓"]
    table(headers, metric_rows(state, "CEM"), [1.4, .95, .95, .95, .95, .95, .95])
    paragraph("areaOracle与原生HRref使用训练视角真实HR，属于额外信息诊断。cam00全程留出，未参与训练或选择本轮CEM端点。表格小数仅用于展示，JSON保留全部原始精度。")
    paragraph("Wu流程为coarse3000+fine14000后再执行SR编号1–6000，继承父模型训练相机颜色affine，并在SR期间再次增密和剪枝。CEM父模型使用作者40k→总50k日程，父阶段没有新增affine，SR固定拓扑。训练预算、容量成长、颜色初始化、优化器和历史选择均不同，跨骨干绝对分数只能比较完整系统，不能单独识别骨干结构的因果影响。")
    paragraph("PSNR为逐帧全图PSNR的算术平均；SSIM使用RGB 11×11高斯窗口、sigma1.5和总体协方差，去除外侧5像素；LPIPS使用AlexNet v0.1。统一评价没有消除训练硬件差异。固定3帧ROI指标和GT相对邻帧差分是独立诊断，不与300帧主均分混合，也不构成几何或时间一致性证明。")

    heading("Wu统一重评与配对增量", page=True)
    table(headers, metric_rows(state, "Wu"), [1.4, .95, .95, .95, .95, .95, .95])
    paragraph("Wu使用已有开发表固定的六个历史检查点，本轮只重渲染和统一重评，不追加训练，也不应用额外posthoc cull、颜色校正或位移场。历史端点的选择影响仍存在，不能称为全新未选择重复。")
    for system in ("CEM", "Wu"):
        heading(system + "系统内链条增量")
        table(["目标−来源", "Cook\nΔPSNR", "Cook\nΔSSIM", "Cook\nΔLPIPS", "Cut\nΔPSNR", "Cut\nΔSSIM", "Cut\nΔLPIPS"], delta_rows(state, system), [1.4, .95, .95, .95, .95, .95, .95])
    paragraph("以上差分均为目标减来源。PSNR和SSIM正值表示改善，LPIPS负值表示改善。两景均值与完整原始差分保存在机器结果中；当前两个开发场景不代表完整标准基准，也不提供显著性结论。")

    heading("数据协议与对照含义", page=True)
    paragraph("本轮输入由HR图像执行4×4面积平均生成，LR为336×252，目标HR为1344×1008。Cook用20台训练相机，Cut缺cam04而使用19台；两景均将cam00的全部300帧留出。旧bicubic LR条件下的CEM指标只作为历史背景，不与本表混合。")
    paragraph("Teacher指冻结的预训练超分网络给训练相机LR生成的伪HR监督图像。原始teacher采用原版SwinIR；area匹配teacher是在area退化上微调后冻结的SwinIR，使其更符合本轮LR输入的形成方式。Teacher不读取留出相机图像，伪HR也不等于真实HR。")
    table(["对照臂", "受控变化"], [[ARM_LABELS[a], state["comparison"]["arm_descriptions"][a]] for a in ARMS], [1.35, 5.75], left=(0, 1))
    paragraph("原生HRref和6k Oracle回答不同问题。原生HRref从与LR父模型相同的LR初始化点云独立开始，以训练HR直接监督作者40k→总50k日程，保留原生拓扑成长过程。Oracle从已完成LR父检查点继续6000步、固定拓扑且无w4，仍用混合目标：第一个采样训练相机的LR项系数为1，另外两台同时间训练相机的HR项各为0.2，并保留原生正则；所有SR臂共有的颜色仿射先作用于渲染。这些系数不等于实测梯度贡献，不能据此断言LR梯度占主导。因此，两者同时改变监督目标、优化历史、拓扑成长和预算，Oracle不是纯HR原生训练，也不是严格重建上界。")
    paragraph("原生HRref通过独立IO包装器调用作者已有的n_job=1接口，按需解码训练HR，避免后台解码结果积压。冻结训练源码、输入像素、相机采样顺序、RAdam和40k→总50k日程保持不变。CTS实际环境的CPU验收确认同种子相机顺序、真实图像逐像素结果和主进程随机状态一致。此工程适配仅用于独立HRref，不改变六个SR臂。")
    paragraph("CEM六个SR臂继续父检查点的原生11组RAdam状态，绝对时钟为50001–56000，共执行6000次参数更新。固定拓扑只限制增密、剪枝和层级结构变化，并未冻结几何、外观、运动或寿命参数。每组实际RAdam步数按有梯度参与记录，完整计数见comparison.json。SR新增训练相机颜色affine从零开始；w4仅在相应臂使用，cam00颜色和位移始终为identity。")

    for scene in SCENES:
        for camera in ("cam00", "cam01"):
            scope = "留出视角" if camera == "cam00" else "训练视角拟合诊断"
            for region in ("whole", "face"):
                item = next(x for x in state["selected"] if x["scene"] == scene and x["camera"] == camera and x["region"] == region and x["role"] == "main_raw")
                for page in report_tile_pages(item):
                    heading(LABELS[scene] + scope + ("完整场景" if region == "whole" else "固定人脸区域") + "第" + str(page["number"]) + "组", page=True)
                    paragraph("所有图保持科学原板顺序。整图用于定位场景结构，人脸采用预先固定ROI。cam01的teacher仅为训练伪HR参照；cam00没有teacher。LR bicubic仅为输入视觉参照。HRref和Oracle使用训练HR，属于额外信息诊断。")
                    report_grid(item, page)
            for text in state["review"]["scene_notes"][scene][camera]:
                paragraph(text)
    paired = [x for x in state["selected"] if x["role"] == "Wu_CEM_paired_raw"]
    for item in paired:
        for page in report_tile_pages(item):
            heading(LABELS[item["scene"]] + "Wu与CEM相同区域第" + str(page["number"]) + "组", page=True)
            paragraph("三列分别为" + "、".join(page["method_columns"]) + "；上行为Wu历史开发端点，下行为CEM固定端点，均为raw渲染。GT使用同一ROI，见本景cam00人脸首组第一幅。完整六臂配对原板保留，未根据结果挑选方法。训练流程和预算差异仍存在，不能据锐度推断几何真实性。")
            report_grid(item, page)
        for text in state["inspected"][item["panel"]["path"]]["observations"]:
            paragraph(text)

    for system in ("CEM", "Wu"):
        heading(system + "训练资源与预算", page=True)
        endpoints = [state["rows"][system, scene, arm] for scene in SCENES for arm in (CEM_ARMS if system == "CEM" else ARMS)]
        table(["场景与端点", "训练GPU", "训练秒数", "峰值GiB", "显存口径", "阶段预算"], [resource_values(e) for e in endpoints], [1.65, 1.25, 1.05, .9, 1.05, 1.2], left=(0, 1, 5))
        if system == "Wu":
            prefix = []
            for scene in SCENES:
                e = state["rows"]["Wu", scene, "areaL0"]
                for phase in e["resources"].get("parent_prefix_stages", []):
                    seconds = phase.get("seconds")
                    peak_gb = phase.get("peak_GB")
                    prefix.append([LABELS[scene], phase.get("stage", "未知"), phase.get("GPU") or "未知", "未知" if seconds is None else f"{seconds:.2f}", "未知" if peak_gb is None else f"{peak_gb * 1e9 / 2**30:.3f}"])
            if prefix:
                table(["场景", "共享前缀", "原始训练GPU", "训练秒数", "allocated GiB"], prefix, [.65, 1., 2.8, 1.35, 1.3], left=(1, 2))
        paragraph("SR秒数为各分支增量，父模型耗时另列并共享，不乘以六臂。原生两阶段训练秒数求和，峰值取各阶段最大值。未知资源明确记为未知，零值仅在原始记录确为零时展示。初始化、教师生成、传输、失败尝试和外部进程成本不包含在训练秒数中。")
        paragraph("CEM列为PyTorch reserved峰值；Wu历史peak_GB为PyTorch allocated字节除以10亿，已换算GiB。二者口径不同，不能直接据此作公平显存排序。Torch allocator记录不覆盖外部Faiss等分配。每项训练设备、阶段时间与原始资源字段均保留在comparison.json中。")
        if system == "CEM":
            failed_seconds = state["io_amendment"]["failed_stage_seconds"]
            paragraph(f"额外失败成本单独保留：Cook旧HRref第一阶段耗时{failed_seconds:.2f}秒（{failed_seconds / 60:.2f}分钟），被CTS内核确认为主机内存不足而终止；这不是GPU显存不足。失败发生于首个30k检查点之前，旧目录和证据保留，未产生可续训的30k或40k检查点。两景HRref按登记的单worker IO适配在新目录从相同LR初始化开始，重启未依据留出指标选择。以上失败耗时单列为额外工程成本，不计入成功端点训练秒数。")
    heading("来源与尚未完成的验收", page=True)
    paragraph("全部28项主值来自已核验result_index或固定Wu清单对应的渲染收据与统一评价summary和300行CSV。检查点SHA、逐帧PNG身份、训练收据、代码和配置身份保存在机器结果中。生成器不读取模型或训练日志，不根据PSNR自动宣布人脸清晰或骨干达到上限。")
    paragraph("本次图像检查为代理对非盲化固定3帧的实际检查，human_acceptance=false，实际人类盲评尚未完成。正文按cam00和cam01的帧150原始tile顺序重排为三列；每页人脸网格最多六幅。科学原板、其他固定帧0和299、训练相机校正后的独立板及全部原始图像保留为补充诊断。Word排版验收单独保存为逐页PNG检查记录；文件生成本身不代表视觉验收。")
    paragraph("下列文件名与SHA256用于定位来源。完整绝对路径和全部机器记录见本报告同名.source.json文件。")
    for label, item in state["sources"].items():
        paragraph(label + "  " + Path(item["path"]).name + "\nSHA256 " + item["sha256"], "Caption")
    for element in (document.element, document.styles.element):
        for border in list(element.xpath(".//w:pBdr")):
            border.getparent().remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.line_spacing = Pt(12)
    footer.add_run("第 ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple");field.set(qn("w:instr"), "PAGE");footer._p.append(field)
    footer.add_run(" 页").font.size = Pt(9)
    document.core_properties.title = title
    document.core_properties.subject = "CEM area4 and fixed Wu endpoints"
    document.core_properties.author = "4dsr research"
    return document, "\n".join(md)


def parser():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--comparison", type=Path, required=True)
    p.add_argument("--diagnostics", type=Path, required=True, help="Formal diagnostic output directory")
    p.add_argument("--review", type=Path, required=True, help="Root actual image observations and scientific interpretation")
    p.add_argument("--protocol-amendment", type=Path, default=BATCH / "protocol_amendment_serial_HR_IO_v2.json", help="Real archived HR serial-IO engineering amendment")
    p.add_argument("--io-cpu-qa", type=Path, default=BATCH / "admission/serial_io_CPU_QA_CTS_v1.json", help="Actual CTS joblib1.6 CPU IO equivalence receipt")
    p.add_argument("--out", type=Path, required=True, help="New docs output stem")
    p.add_argument("--check-only", action="store_true", help="Read-only stdlib preflight; no docx import/output")
    return p


def main():
    args = parser().parse_args()
    state = preflight(args)
    if args.check_only:
        print(json.dumps(dict(status="complete_inputs_and_actual_review_verified_no_output", endpoints=28, reviewed_original_panels=len(state["selected"]), report_layout="three_columns_native_10pt_labels", docx_imported="docx" in sys.modules, output_created=False)))
        return
    document, markdown = author(state)
    stem = state["stem"]
    stem.parent.mkdir(parents=True, exist_ok=True)
    docx_path, md_path, receipt_path = [Path(str(stem) + suffix) for suffix in (".docx", ".md", ".source.json")]
    with docx_path.open("xb") as stream:
        document.save(stream)
    with md_path.open("x") as stream:
        stream.write(markdown)
    receipt = dict(schema="CEM_area4_Word_source_receipt_v2", status="created_pending_root_document_render_QA", created_utc=datetime.now(timezone.utc).isoformat(),
        sources=state["sources"], DOCX=identity(docx_path), Markdown=identity(md_path), endpoints=28,
        metric_sources=[dict(system=e["system"], scene=e["scene"], arm=e["arm"], original_values=e["metrics"]["values"],
            summary=e["metrics"]["summary"], per_frame=e["metrics"]["per_frame"], checkpoint=e["checkpoint"],
            budget=e["budget"], resources=e["resources"]) for e in state["comparison"]["endpoints"]],
        HR_serial_IO_amendment=state["io_amendment"],
        formatted_metric_tables={system: metric_rows(state, system) for system in ("CEM", "Wu")},
        reviewed_original_panels=[dict(path=x["panel"]["path"], sha256=x["panel"]["sha256"], role=x["role"], reviewer=state["inspected"][x["panel"]["path"]]["reviewer"], original_panel_embedded=False) for x in state["selected"]],
        inserted_report_tiles=state["report_image_ledger"],
        report_layout=dict(columns=3, labels="native Word 10pt", face_tiles_per_page_max=6, whole_tiles_per_page_max=12,
                           original_scientific_diagnostics_modified=False, source_pixels_enhanced=False,
                           derived_display_assets=str(Path(str(stem) + ".report_assets"))),
        human_acceptance=False, rendered_page_images_inspected=False, scientific_interpretation_source=state["sources"]["actual_review"],
        marker_and_render_operations="external caller responsibility; never executed by this generator", bundled_python=str(Path(sys.executable).resolve()))
    with receipt_path.open("x") as stream:
        json.dump(receipt, stream, indent=2, ensure_ascii=False, allow_nan=False);stream.write("\n")
    print(json.dumps(dict(status=receipt["status"], DOCX=receipt["DOCX"], Markdown=receipt["Markdown"], source_receipt=identity(receipt_path)), ensure_ascii=False))


if __name__ == "__main__":
    main()
