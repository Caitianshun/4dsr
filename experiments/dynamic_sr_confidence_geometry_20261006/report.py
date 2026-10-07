"""Rebuild the local Chinese phase report from frozen machine-readable evidence.

Run with the bundled document Python runtime, after the document operation
marker required by the documents skill. This generator does not publish, train,
alter evaluation weights, or claim that pending endpoints have finished.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
from pathlib import Path
import statistics

ROOT = Path(__file__).resolve().parents[2]
RUN = ROOT/'output/dynamic_sr_confidence_geometry_20261006'
DEFAULT_STEM = ROOT/'docs/dynamic_sr_confidence_geometry_results_2026-10-06'
CORE = ('C1', 'Jperm', 'B2perm', 'R', 'G', 'RG')


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(1048576), b''):
            h.update(block)
    return h.hexdigest()


def read(path, default=None):
    p = Path(path)
    return json.loads(p.read_text()) if p.exists() else default


def rows(path):
    p = Path(path)
    if not p.exists():
        return []
    with p.open() as f:
        return list(csv.DictReader(f))


def fmt(value, digits=6):
    return '待完成' if value is None else f'{float(value):.{digits}f}'


def snapshot():
    files = dict(summary=RUN/'quality_summary.json', calibration=RUN/'calibration.json',
        caches=RUN/'cache_manifest.json', confidence=RUN/'cache/confidence_view/cache_manifest.json',
        operator=RUN/'operator_checks_cuda3090_full.json', geometry=RUN/'geometry_cuda_checks_v2.json',
        cause=RUN/'cause_matrix_v1/cause_matrix.json', source_incident=RUN/'source_incidents/incident.json',
        calibration_failure=RUN/'calibration_failed.json', backward_diagnostic=RUN/'backward_equivalence_failure_diagnostic.json',
        pending_coefficients=RUN/'calibration_coefficients_pending_shadow.json',
        protocol=RUN/'protocol.json', state=RUN/'state.json')
    for archive in sorted(RUN.glob('calibration_failure_*/calibration_failed.json')):
        files['failure_'+archive.parent.name] = archive
    result = {k:read(p, {}) for k, p in files.items()}
    result['source_identity'] = {k:dict(path=str(p.relative_to(ROOT)), sha256=sha(p)) for k, p in files.items() if p.exists()}
    for name in ('quality_summary','quality_per_endpoint','cost_ledger','confidence_comparison'):
        path=RUN/(name+'.csv')
        if path.exists():result['source_identity'][name+'_csv']=dict(path=str(path.relative_to(ROOT)),sha256=sha(path))
    result['calibration_failures'] = [result[key] for key in files if key.startswith('failure_')]
    if result['calibration_failure']:result['calibration_failures'].append(result['calibration_failure'])
    result['quality'] = rows(RUN/'quality_summary.csv')
    result['endpoint_quality'] = rows(RUN/'quality_per_endpoint.csv')
    result['costs'] = rows(RUN/'cost_ledger.csv')
    result['comparison'] = rows(RUN/'confidence_comparison.csv')
    return result


def report_content(data):
    """One content tree produces Word and its traceable Markdown source."""
    summary = data['summary']; execution = summary.get('execution', {})
    completed = execution.get('completed_primary_endpoints', 0)
    completed_core = execution.get('core_evidence_complete', False)
    phase = '核心证据已汇总' if completed_core else '训练与工程核验进行中'
    sections = []
    def heading(value, level=1):
        sections.append(dict(type='heading', text=value, level=level))
    def paragraph(value):
        sections.append(dict(type='paragraph', text=value))
    def table(headers, values, widths):
        sections.append(dict(type='table', headers=headers, rows=values, widths=widths))
    def page():
        sections.append(dict(type='page'))
    sections.append(dict(type='title', text='动态场景超分置信度与软几何阶段报告'))
    paragraph('2026年10月6日  cook_spinach 完整动态场景四倍超分  状态 '+phase)
    paragraph(f'本轮验证外部超分（SR）细节是否通过共享动态表示妨碍真实低分辨率（LR）观测的拟合，以及置信度选择和弱深度约束能否改善新视角。核心计划为六臂各两套续训，共12个端点。目前主评价完成{completed}/12；额外浮点评价、原因干预和模块校准分别按来源记录列示。未完成的证据不用于宣布方法有效或安排新的方法组合。')
    heading('本轮解决的问题')
    paragraph('同一时刻的训练图像来自不同相机。某个视角中看起来合理的教师纹理，在其他视角或时间可能不一致；共享高斯与形变网络会把这种监督传到其他观察。墙面大尺度误差也可能来自几何、覆盖或外观，不能仅凭图像锐度认定超分成功。我们在同一父状态上控制观察安排、细节监督和几何约束，逐项检验这些解释。')
    table(['分支', '每步观察与监督', '验证作用'], [
        ['C1', '同一HR渲染接真实LR与完整SR监督', '同图母体'],
        ['Jperm', 'LR与置换SR各渲染一次 累积后更新', '观察耦合与顺序'],
        ['B2perm', 'LR加两张置换SR SR系数各0.05', '同加权曝光下平均梯度'],
        ['R', '真实LR加置信加权零空间细节', '观测与补充细节分工'],
        ['G', 'C1加真实LR相对深度排序', '弱几何的受控作用'],
        ['RG', '同图母体同时使用R与G', '互补或冲突']], [1.7, 9.3, 6.4])
    heading('固定协议与预算')
    paragraph('合法输入为19个训练相机cam02至cam20、60个偶数帧，共1140观察。LR为336×252，高分辨率（HR）画布为1344×1008。共同U6000恢复模型、两个Adam优化器及随机状态，固定132972点；每套追加6000次更新，3000仅作恢复，6000为唯一主端点。两张观察表来自预登记局部种子2026100601和2026100602。')
    paragraph('核心预算72000次更新，另有三条各500步LR-only原因探针。条件性动态证据与组件消融60000步，S/V备选24000步，需核心质量复核后决定。失败重算和辅助前向单列。短窗及两个共享父状态后缀都不能替代独立完整场景验证。')
    page()
    heading('监督分工的机制与可验证边界')
    paragraph('R把真实LR能够约束的分量交给观测损失；辅助教师只约束实际线性降采样无法看到的图像变化。实际退化采用bicubic、抗混叠和半像素边界约定；线性部分的伴随与伪逆用于构造零空间投影Q。训练先投影渲染与教师的差，再在投影后的绝对误差上乘冻结空间权重。位置次序不能交换，也不能用普通放大算子代替伴随。')
    paragraph('这个结构可以使辅助损失进入渲染图像的梯度处于线性退化的零空间。但图像梯度随后经过渲染器对共享参数的导数，Adam又使用历史动量；实际参数更新仍可能改变LR输出、其他相机和其他时间。因此需要32个固定合法探针及真实Adam单步影子更新，不能把算子恒等式当作“LR一定不变”的保证。')
    paragraph('G只使用真实训练LR产生的Depth Anything V2相对逆深度。辅助渲染输出覆盖权重alpha、深度加权的一阶矩和二阶矩，在HR画布上先以非负area权重聚合到LR，再除以alpha获得平均深度。局部点对排序损失更新有效位置；协方差和不透明度在该分支停止梯度。训练掩码取冻结父模型支持，防止模型通过降低当前alpha逃避监督。')
    paragraph('排序约束说明哪个位置更近，不提供米制深度真值；混合多个表面的射线平均深度也不等于唯一表面。稳定归一化和近似平局带降低噪声影响。其权重只在固定32观察上按真实LR位置梯度量级校准，前500步渐入，不依据开发HR搜索权重。')
    heading('冻结置信度与零更新诊断')
    conf = data['confidence']; info = conf.get('summary', {})
    paragraph(f'全部1140置信图已离线冻结。教师回到真实LR的局部误差形成闭环证据；同刻至多两个近邻训练相机提供空间证据。对应必须满足正深度、视野、往返一致、遮挡及真实LR Census和光度支持，并以投影足迹抗混叠比较教师残差。低纹理或无有效邻居标为未知，回退0.5。当前平均权重{fmt(info.get("mean_weight"))}，Kish有效比例{fmt(info.get("kish_effective_fraction"))}，有效对应{fmt(info.get("mean_valid_fraction"),4)}，未知{fmt(info.get("mean_unknown_fraction"),4)}。Kish比例描述权重集中度，不是真实正确率。')
    comparison = []
    labels = dict(old_T='旧T需求图', CLEAR_style='CLEAR风格闭环', IE_style='IE风格分歧', new='本轮冻结置信')
    for method in ('old_T', 'CLEAR_style', 'IE_style', 'new'):
        rr = [r for r in data['comparison'] if r['method'] == method and r['scope'] == 'full']
        if rr:
            mean_weight = statistics.mean(float(r['mean_weight']) for r in rr)
            mass = sum(float(r['mean_weight']) for r in rr)
            weighted = sum(float(r['mean_weight'])*float(r['weighted_teacher_abs_error']) for r in rr)/mass
            rho = statistics.mean(float(r['spearman_abs_error']) for r in rr if r['spearman_abs_error'])
            comparison.append([labels[method], fmt(mean_weight), fmt(weighted,8), fmt(rho,5)])
        else:
            comparison.append([labels[method], '待完成', '待完成', '待完成'])
    table(['固定train76诊断图', '平均权重', '加权教师绝对误差', '逐观察相关均值'], comparison, [5.4, 3.0, 4.8, 4.2])
    paragraph('本轮权重选择的加权误差较低，但误差相关性并未优于IE风格分歧；较低权重也明显削弱教师监督。上述诊断不能证明训练收益。CLEAR和IE仅为公式诊断适配，未声称完整复现。未知、局部对比度、LR动态代理、边界和置信区间分层完整保留；HR只由隔离评价程序读入，不改变置信缓存或阈值。')
    page()
    heading('工程核验与保留的异常证据')
    operator = data['operator']; cases = operator.get('cases', [])
    full = next((c for c in cases if c.get('operator', {}).get('hr_size') == [1008,1344] and c['operator'].get('runtime_dtype') == 'torch.float32'), {})
    measured = full.get('measurements', {})
    geom = data['geometry'].get('gpu', {})
    finite = geom.get('finite_differences', [])
    max_fd = max((r['absolute_error'] for r in finite), default=None)
    calibration = data['calibration']; passed_calibration = calibration.get('status') == 'passed'
    table(['检查', '当前记录', '支持范围'], [
        ['真实退化与Q', '完整CPU与3090算子检查通过' if cases and all(c.get('passed') for c in cases) else '待完成', '退化 伴随 幂等 反向泄漏'],
        ['G CUDA稳定窗口', '通过 最大有限差分误差 '+fmt(max_fd,8) if geom.get('status') == 'passed' else '待完成', '合成矩与位置梯度路径'],
        ['32观察校准', '通过' if passed_calibration else '待工程复核', 'k_R lambda_G及Adam影子更新']], [4.5, 7.0, 5.9])
    if measured:
        paragraph('完整画布float32算子核验记录了真实降采样一致性、Q幂等和细节梯度泄漏；具体数值及登记容差保存在operator_checks_cuda3090_full.json。几何CUDA检查中协方差与不透明度梯度为空，位置梯度非零。以上是实现正确性的有限核验，不能替代父模型上的校准或质量结果。')
    incident = data['source_incident']
    if incident:
        paragraph('缓存生成启动后，depth_prior.py曾被追加来源和验收字段，导致进程实际加载的旧函数与收据末尾记录的文件哈希可能不一致。已重建并核对实际初始源码，矩渲染、area聚合、DAV2推理和排序算法未改变。实际加载源码SHA前12位为'+incident.get('actual_loaded_source_sha256','')[:12]+'；后续文件SHA前12位为'+incident.get('later_file_sha256','')[:12]+'。原收据和差异补丁保留，另以来源关联记录纠正身份，不能静默改写原记录。')
    if geom.get('status') == 'passed':
        paragraph('首次G有限差分测试把空射线纳入窗口，微扰改变栅格支持，使归一深度在分母下界附近跳变。修复测试样例后使用支持稳定的7至9窗口，alpha最小值约0.3975，测试窗口内HR支持变化为0；原失败及分析仍保留，未放宽登记容差，也未修改训练损失。这说明原失败受测试条件影响，不说明真实场景中所有遮挡边缘都可微。')
    if data['calibration_failures']:
        count=len(data['calibration_failures'])
        if passed_calibration:
            paragraph(f'32观察校准保留{count}次失败归档，涉及联合与分开反向及完整Adam单步等价检查。最新通过记录中的重复计算基线、登记容差和完成状态才决定可派发模块；失败耗时不算作正式训练更新。')
        else:
            paragraph(f'32观察校准保留{count}次失败归档，涉及联合与分开反向及完整Adam单步等价检查，正式训练更新均为0。重复同一路径显示部分差异与CUDA数值波动同量级，但不能据此自动通过验收。固定梯度计算已有候选系数，仍待Adam影子更新正式通过；R/G/RG派发不使用未验收系数。')
    heading('原因判断仍有多个未决因素')
    causes = data['cause'].get('rows', [])
    counts = {status:sum(r['status'] == status for r in causes) for status in ('已确认实现问题','有干预支持的贡献因素','相关现象','证据不足')}
    paragraph('当前原因矩阵按独立证据标注：'+ '，'.join(status+str(n)+'项' for status,n in counts.items())+'。开发LR颜色修正未稳定改善cam01；固定二维位移可减小部分开发残差，但可能补偿错误几何。它们使用目标开发LR，只是原因诊断，不进入合法新视角主成绩。')
    paragraph('同一表面的真实训练覆盖、方向外观、动态背景污染及P_joint、P_xyz、P_SH与C1第500步的受控干预尚需补齐。没有足够证据时保留“尚不能确定”，不选择单一最顺眼的原因。')
    page()
    heading('质量结果与当前完成度')
    paragraph('主指标先逐帧平均，再对cam00与cam01等权，最后对两个后缀等权。PSNR评价像素误差，SSIM评价结构相似，LPIPS评价感知距离；三者分别报告，不拼接综合分。FFT报告low、mid、high绝对MSE和比例；区域、时间与矩支持用于解释取舍。渲染和指标统一在本机3090复算，训练硬件差异仍可能影响优化。')
    quality = []
    for arm in ('U6000','J1','Async2','Sync2','Repeat7','Video7',*CORE):
        candidate = next((r for r in data['quality'] if r['arm'] == arm and r['scope'] == 'equal_camera_mean'), None)
        if candidate:
            quality.append([arm, fmt(candidate['psnr'],5), fmt(candidate['ssim']), fmt(candidate['lpips']),
                            '历史' if arm not in CORE else ('两套完整' if candidate['status'] == 'complete' else '仅'+candidate['repeat_count']+'套')])
        else:
            quality.append([arm, '待完成','待完成','待完成','待完成'])
    table(['方法', 'PSNR', 'SSIM', 'LPIPS', '状态'], quality, [3.5, 3.3, 3.3, 3.3, 4.0])
    if not completed_core:
        paragraph('核心结果尚未齐备，表中单套数值仅为运行状态记录，不能与历史两套均值构成完整方法结论。尤其RG只优于较弱C1时，也不能据此宣称超过当前最佳。两个后缀共享U6000，仅检验续训重复性，不是两个独立从头训练种子。')
    else:
        paragraph('核心两套证据已齐备，应结合配对差值、区域与时序取舍决定是否进入空间、时间、联合证据及权重总量对照。若三指标均稳定退步且无实际价值，应保留负结果并收缩路线，不通过新增模块维持原叙事。')
    heading('实际资源与后续任务')
    costs = data['costs']
    if costs:
        compact = [[r['endpoint'],r['training_gpu'].replace('NVIDIA ',''),fmt(r['training_seconds'],1),fmt(r['peak_gb'],2)] for r in costs]
        table(['已记录端点', '训练硬件', '训练秒', '峰值GB'], compact, [3.9,7.1,3.2,3.2])
    else:
        paragraph('正式端点的训练秒数和峰值显存待完成收据；资源安排为第一套本机PRO6000、第二套A100，实际执行以每分支完成记录为准。A100另一张忙卡保持原任务，cts连接超时未派发，5090不用于训练。')
    paragraph('后缀1配置为本机PRO6000，后缀2为A100，实际硬件以各分支收据为准。先完成核心、校准及原因探针，训练返回即接评价、回传与完整性核验；每小时监测作异常兜底。额外评价保存196个浮点RGB及area矩，统一GPU1锁。条件性分支等完整质量复核后登记。')
    paragraph('可追踪来源位于本轮output目录的protocol.json、quality_summary.json、calibration.json及失败记录、cache_manifest.json、operator_checks_cuda3090_full.json、geometry_cuda_checks_v2.json、source_incidents和cause_matrix_v1。逐帧表、ROI、FFT、成本和完整检查点用于后续复算；报告来源SHA完整保存于同目录report_snapshot.json。')
    return sections, phase


def save_markdown(path, blocks):
    content = []
    for block in blocks:
        kind = block['type']
        if kind == 'title':
            content += ['# '+block['text'], '']
        elif kind == 'heading':
            content += ['#'*(block['level']+1)+' '+block['text'], '']
        elif kind == 'paragraph':
            content += [block['text'], '']
        elif kind == 'table':
            content += ['| '+' | '.join(block['headers'])+' |', '| '+' | '.join(['---']*len(block['headers']))+' |']
            content += ['| '+' | '.join(str(v) for v in row)+' |' for row in block['rows']]
            content.append('')
    path.write_text('\n'.join(content)+'\n')


def save_docx(path, blocks, figures=None):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
    document = Document()
    section = document.sections[0]; section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.left_margin = section.right_margin = Cm(1.8); section.top_margin = section.bottom_margin = Cm(1.6)
    # Word's bundled template includes a CJK grid and a coloured Title border.
    # Explicit leading avoids font-metric-dependent double line heights in LO.
    for grid in list(section._sectPr.findall(qn('w:docGrid'))):
        section._sectPr.remove(grid)
    for name, size in [('Normal',11),('Title',23),('Heading 1',15),('Heading 2',12)]:
        style = document.styles[name]; style.font.name = 'Noto Sans CJK SC'; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0,0,0)
        rpr=style.element.get_or_add_rPr();fonts=rpr.get_or_add_rFonts()
        fonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            fonts.attrib.pop(qn('w:'+attr),None)
        cs=rpr.find(qn('w:szCs'))
        if cs is None:cs=OxmlElement('w:szCs');rpr.append(cs)
        cs.set(qn('w:val'),str(size*2))
        ppr=style.element.get_or_add_pPr()
        for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');ppr.append(snap)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = Pt(15.5 if name=='Normal' else size+5)
    document.styles['Normal'].paragraph_format.first_line_indent = Cm(0)
    document.styles['Heading 1'].paragraph_format.space_before = Pt(8)
    document.styles['Heading 1'].paragraph_format.keep_with_next = True
    def font_run(run, size=11, bold=False):
        run.font.name='Noto Sans CJK SC'; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=RGBColor(0,0,0)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
    for block in blocks:
        kind=block['type']
        if kind == 'title':
            p=document.add_paragraph(block['text'],'Title');p.paragraph_format.space_after=Pt(9)
        elif kind == 'heading':
            document.add_paragraph(block['text'],'Heading '+str(block['level']))
        elif kind == 'paragraph':
            document.add_paragraph(block['text'])
        elif kind == 'page':
            document.add_page_break()
        elif kind == 'table':
            table=document.add_table(rows=1,cols=len(block['headers']));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
            for column,width in zip(table.columns,block['widths']):
                column.width=Cm(width)
            borders=OxmlElement('w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                element=OxmlElement('w:'+side);element.set(qn('w:val'),'single');element.set(qn('w:sz'),'4');element.set(qn('w:color'),'D9D9D9');borders.append(element)
            table._tbl.tblPr.append(borders)
            for i, header in enumerate(block['headers']):
                table.rows[0].cells[i].text=header
            head=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(head)
            for row in block['rows']:
                cells=table.add_row().cells
                for cell,value in zip(cells,row):cell.text=str(value)
            for ridx,row in enumerate(table.rows):
                no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
                for cidx,cell in enumerate(row.cells):
                    cell.width=Cm(block['widths'][cidx]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    props=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
                    for side in ('top','left','bottom','right'):
                        el=OxmlElement('w:'+side);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');margin.append(el)
                    props.append(margin)
                    if ridx==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9F1F7');props.append(shade)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(14)
                        p.alignment=WD_ALIGN_PARAGRAPH.LEFT if cidx==0 or (len(block['headers'])==3 and cidx==1) else WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs:font_run(run,size=10,bold=ridx==0)
            document.add_paragraph().paragraph_format.space_after=Pt(1)
    if figures:
        document.add_page_break();document.add_paragraph('固定同帧图像对照','Heading 1')
        for figure in figures:
            image_path=Path(figure['path'])
            if not image_path.is_absolute():image_path=ROOT/image_path
            if not image_path.exists():raise ValueError('Registered figure missing '+str(image_path))
            document.add_paragraph(figure['caption'])
            document.add_picture(str(image_path),width=Cm(17.4))
    document.core_properties.title='动态场景超分置信度与软几何阶段报告'
    document.core_properties.subject='固定短窗的受控训练与验证'
    document.core_properties.author='4dsr research'
    document.save(path)


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--stem',type=Path,default=DEFAULT_STEM)
    parser.add_argument('--figures-index',type=Path);args=parser.parse_args()
    data=snapshot();blocks,phase=report_content(data);args.stem.parent.mkdir(parents=True,exist_ok=True)
    figures=None
    if args.figures_index:
        if not data['summary'].get('execution',{}).get('core_evidence_complete'):
            raise ValueError('Final figure insertion requires complete core evidence')
        registered=read(args.figures_index);figures=registered['figures']
        assert all(int(f['frame']) in (40,80) and f['camera'] in ('cam00','cam01','cam02') for f in figures)
    save_markdown(args.stem.with_suffix('.md'),blocks)
    save_docx(args.stem.with_suffix('.docx'),blocks,figures)
    receipt=dict(status=phase,source_identity=data['source_identity'],source_sha256=sha(__file__),
                 docx_path=str(args.stem.with_suffix('.docx')),docx_sha256=sha(args.stem.with_suffix('.docx')),
                 md_path=str(args.stem.with_suffix('.md')),md_sha256=sha(args.stem.with_suffix('.md')),
                 quality_execution=data['summary'].get('execution'),visual_QA='pending_render_and_each_page_inspection')
    (RUN/'report_snapshot.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(receipt,ensure_ascii=False),flush=True)


if __name__ == '__main__':
    main()
