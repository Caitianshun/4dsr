"""Build a final core report only after the frozen evidence has completed.

Independent wrapper around report.py's verified 11pt document layout. No
training, GPU work or change to synchronized sources is performed. Use the
bundled document Python. The existing DOCX operation marker is not repeated.
--check-ready reads state only; the default build refuses incomplete evidence.
"""
from __future__ import annotations
import argparse
from collections import Counter
import datetime
import json
from pathlib import Path
import statistics
import report as stage

ROOT, RUN=stage.ROOT,stage.RUN
ROI=ROOT/'output/dynamic_sr_prior_diagnosis_20260929/spectrum/roi_protocol.json'
DATA=ROOT/'data/dynamic_sr/n3dv_prepared/cook_spinach'
FIGURE_SELECTION=(('cam00','hand_utensil_pan_reference'),('cam01','lamp_wall_corner_reference'))
PANEL_LABELS=('HR参考','U6000父模型','C1 后缀1','RG 后缀1')
PREPROCESSING_SHA256='61871fe88a05e3bcc6a4a84a481448e2ada9d787f9038f75ab1747eb5ee2d492'
FUTURE_ASSETS_SHA256='7c657e7590f5917e9a337232d93ecafba128bb976f6d131959ae04fba139d515'


def readiness():
    reasons=[]
    summary=stage.read(RUN/'quality_summary.json',{})
    state=stage.read(RUN/'core_finalization_state.json',{})
    integrity=stage.read(RUN/'final_integrity.json',{})
    if not summary.get('execution',{}).get('core_evidence_complete'):reasons.append('Core quality/extra evidence incomplete')
    if state.get('status')!='core_complete_pending_research_and_visual_review':reasons.append('Core finalization completion event absent')
    if integrity.get('status')!='passed_core_endpoint_evidence' or len(integrity.get('endpoints',[]))!=12:reasons.append('Twelve endpoint integrity receipt absent')
    if stage.read(RUN/'calibration.json',{}).get('status')!='passed':reasons.append('Calibration not passed')
    for path in [RUN/'cause_interpretation.json',RUN/'P1_recovery_matrix_v1/cause_matrix.json',RUN/'P1_recovery_matrix_v1/complete.json',
                 RUN/'calibration_total_cost.json',RUN/'preprocessing_cost_ledger.json',RUN/'future_benchmark_readiness.json',RUN/'phase13_decision.json',
                 RUN/'accounting_final_core/execution_index.json',RUN/'figures/index.json',ROI]:
        if not path.exists():reasons.append('Missing '+str(path.relative_to(ROOT)))
    if (RUN/'future_benchmark_readiness.json').exists() and stage.sha(RUN/'future_benchmark_readiness.json')!=FUTURE_ASSETS_SHA256:
        reasons.append('Frozen future-asset audit identity changed')
    if not stage.read(RUN/'phase13_decision.json',{}):reasons.append('Root quality judgment and section 13 decision absent')
    if (RUN/'preprocessing_cost_ledger.json').exists():
        try:verify_preprocessing(stage.read(RUN/'preprocessing_cost_ledger.json',{}))
        except (AssertionError,KeyError,OSError,ValueError) as error:reasons.append('Preprocessing cost identity invalid: '+str(error))
    return dict(status='ready_for_final_core_report' if not reasons else 'waiting_core_completion',reasons=reasons,
                completed_primary=summary.get('execution',{}).get('completed_primary_endpoints',0))


def snapshot():
    data=stage.snapshot()
    paths=dict(cause_interpretation=RUN/'cause_interpretation.json',cause=RUN/'P1_recovery_matrix_v1/cause_matrix.json',
        P1_recovery=RUN/'P1_recovery_matrix_v1/complete.json',calibration_total=RUN/'calibration_total_cost.json',
        preprocessing=RUN/'preprocessing_cost_ledger.json',future_assets=RUN/'future_benchmark_readiness.json',
        phase13=RUN/'phase13_decision.json',integrity=RUN/'final_integrity.json',finalization=RUN/'core_finalization_state.json',
        accounting=RUN/'accounting_final_core/execution_index.json',gallery=RUN/'figures/index.json',roi=ROI)
    for candidate in (RUN/'phase13/state.json',RUN/'phase13_dispatch_state.json'):
        if candidate.exists():paths['phase13_execution']=candidate;break
    for key,path in paths.items():
        data[key]=stage.read(path,{})
        if path.exists():data['source_identity'][key]=dict(path=str(path.relative_to(ROOT)),sha256=stage.sha(path))
    data['source_identity']['stage_generator']=dict(path=str(Path(stage.__file__).relative_to(ROOT)),sha256=stage.sha(stage.__file__))
    for name in ('frequency_summary','frequency_budget','regional_quality'):
        path=RUN/(name+'.csv')
        if path.exists():data['source_identity'][name+'_csv']=dict(path=str(path.relative_to(ROOT)),sha256=stage.sha(path))
    return data


def require_bound(item):
    path=Path(item['path']);path=path if path.is_absolute() else ROOT/path
    if stage.sha(path)!=item['sha256']:raise ValueError('Evidence bytes changed '+str(path))
    return path


def verify_preprocessing(ledger):
    """Verify receipt identities and non-additive accounting without any GPU work."""
    assert stage.sha(RUN/'preprocessing_cost_ledger.json')==PREPROCESSING_SHA256,'Frozen preprocessing ledger changed'
    assert ledger['status']=='completed_receipt_accounting_with_explicit_unknowns'
    refs=[e for entry in ledger['stages'] for e in entry['evidence']]
    refs.extend(ledger['referenced_separate_accounts'].values())
    assert len(refs)==ledger['validation']['bound_evidence_reference_count']==76,'Expected 75 stage references and one separate calibration reference'
    for item in refs:require_bound(item)
    assert len(ledger['stages'])==27 and len({entry['id'] for entry in ledger['stages']})==27
    cal=ledger['referenced_separate_accounts']['calibration'];require_bound(cal)
    assert cal['add_again_to_this_ledger'] is False
    assert cal['sha256']==ledger['summary']['calibration_reference_only_not_in_any_summary_sum']
    assert ledger['validation']['contained_intervals_excluded_from_sum'] is True
    assert ledger['validation']['calibration_not_in_any_stage_or_summary_sum'] is True
    assert ledger['summary']['sum_is_not_complete_wall_clock_or_parallel_makespan'] is True
    assert ledger['summary']['all_scene_formal_training_updates']==0
    return dict(bound_evidence_references=len(refs),stages=len(ledger['stages']),formal_preprocessing_updates=0,
                calibration_counted_again=False,complete_wall_clock='unknown')


def fmt(x,n=6):return stage.fmt(x,n)
def sci(x):return '待记录' if x is None else f'{float(x):.6g}'
def signed(x,n=6):return '待记录' if x is None else f'{float(x):+.{n}f}'


def content(data):
    """Five text pages, followed by two preregistered native-pixel comparisons."""
    blocks=[]
    def p(text):blocks.append(dict(type='paragraph',text=text))
    def h(text):blocks.append(dict(type='heading',text=text,level=1))
    def table(headers,rows,widths):blocks.append(dict(type='table',headers=headers,rows=rows,widths=widths))
    def page():blocks.append(dict(type='page'))
    summary=data['summary'];means=summary['means'];cal=data['calibration'];cost=data['calibration_total']
    blocks.append(dict(type='title',text='动态场景超分置信度与软几何核心结果'))
    p('2026年10月6日执行协议  cook_spinach完整场景短窗四倍超分  核心六臂两套已完成')
    p('本轮检验超分辨率（SR）教师细节是否通过共享动态表示影响真实低分辨率（LR）观测。我们比较置信选择与软几何约束对新视角的作用。高分辨率（HR）只用于评价。12个追加6000步端点、主评价和每端点196观察的浮点评价均通过完整性核验。完成范围是固定开发短窗的受控比较，独立完整场景仍待验证。')
    h('统一质量结果')
    p('PSNR是峰值信噪比，衡量像素误差；SSIM是结构相似度；LPIPS是学习式感知图像块相似度，用预训练特征比较感知差异。前两者越高越好，LPIPS越低越好。')
    q=[]
    for arm in ('U6000','J1','Async2','Sync2','Repeat7','Video7',*stage.CORE):
        row=next(r for r in data['quality'] if r['arm']==arm and r['scope']=='equal_camera_mean')
        q.append([arm,fmt(row['psnr'],5),fmt(row['ssim']),fmt(row['lpips']),'核心两套' if arm in stage.CORE else '历史复用'])
    table(['方法','PSNR','SSIM','LPIPS','来源'],q,[3.3,3.4,3.4,3.4,3.9])
    pair=summary['paired']['RG-minus-C1']
    p('RG相对C1的配对均值变化：PSNR '+signed(pair['delta']['psnr'],5)+' dB，SSIM '+signed(pair['delta']['ssim'])+'，LPIPS '+signed(pair['delta']['lpips'])+'。逐后缀差值与方向一致性保存在quality_summary.json。相关帧不作为独立样本做显著性检验，三指标不合并为综合分。')
    best=[min(means,key=lambda m:means[m][metric]) if metric=='lpips' else max(means,key=lambda m:means[m][metric]) for metric in ('psnr','ssim','lpips')]
    p('本表三指标各自最佳依次为'+ '、'.join(best)+'。最终取舍还需结合频带、区域、时序与成本；只超过较弱C1不能据此声称超过当前最佳。cam00/01已经用于开发，两套共享同一U6000，只检验续训重复性。训练硬件以配置和完成收据为准。统一3090评价不消除硬件对优化的影响。')
    page();h('监督分工与冻结置信度')
    p('同一时刻的教师纹理可能随相机变化。共享高斯和形变网络会把局部监督传播到其他视角与时间，因此表面锐利不自动表示真实细节。C1在同一HR渲染上连接真实LR与完整SR；Jperm将LR与置换SR分别渲染后累积更新；B2perm再平均两张置换SR梯度，每张系数0.05。六臂保持相同更新数与登记观察曝光。')
    p('R将实际双三次（bicubic）抗混叠降采样的线性部分记为D0。D0的伴随把LR梯度回传到HR网格，伪逆估计LR可见分量。零空间投影Q去除该分量。先投影渲染与教师的差，再在投影后的绝对误差上乘冻结空间权重，系数为0.1*k_R；权重与投影不能交换，普通图像放大不是伴随。Q中的变化在图像域不改变D0，但通过渲染器导数和继承动量的Adam以后，参数更新仍可能改变真实LR、邻视角及邻时间。')
    p('G只对真实训练LR使用预训练深度模型Depth Anything V2，获得相对逆深度。辅助渲染先把覆盖权重alpha及深度一阶、二阶矩按非负area权重聚合到LR，再除以alpha获得射线平均深度。固定局部点对约束远近排序，协方差和不透明度在此分支停止梯度；掩码使用冻结父支持，防止降低当前alpha逃避监督。这里只停止辅助分支的直接梯度。共享形变网络更新后，仍可间接改变有效外观、尺度和不透明度。排序不提供米制深度真值，多表面混合的均值也不等于唯一真实表面。RG在同一渲染图像上同时使用R和G。')
    conf=data['confidence']['summary']
    p('置信缓存只使用cam02至cam20的60偶数帧，共1140合法观察。真实LR闭环误差与同刻至多两个近邻训练相机提供证据；正深度、视野、往返、遮挡以及真实LR局部亮度排序（Census）和光度支持共同判定对应有效。投影后的像素覆盖范围决定抗混叠粗尺度。低纹理和全部对应缺失保留未知，回退0.5。平均权重'+fmt(conf['mean_weight'])+'，有效对应比例'+fmt(conf['mean_valid_fraction'],4)+'，未知'+fmt(conf['mean_unknown_fraction'],4)+'；Kish有效样本比例'+fmt(conf['kish_effective_fraction'])+'只描述权重集中度。')
    compared=[]
    names={'old_T':'旧T需求图','CLEAR_style':'CLEAR风格闭环','IE_style':'IE风格分歧','new':'本轮冻结置信'}
    for method,label in names.items():
        rs=[r for r in data['comparison'] if r['method']==method and r['scope']=='full'];assert len(rs)==76
        mass=sum(float(r['mean_weight']) for r in rs)
        weighted=sum(float(r['mean_weight'])*float(r['weighted_teacher_abs_error']) for r in rs)/mass
        compared.append([label,fmt(statistics.mean(float(r['mean_weight']) for r in rs)),fmt(weighted,8),fmt(statistics.mean(float(r['spearman_abs_error']) for r in rs),5)])
    table(['固定train76诊断','平均权重','加权教师绝对误差','相关均值'],compared,[5.4,3.0,5.0,4.0])
    p('本轮权重的加权教师误差较低，但相关性未优于IE风格；明显降权也会减弱教师监督。这是零更新、HR隔离的选择诊断，不能替代训练收益。CLEAR与IE仅为公式适配，未称完整原方法复现。所有阈值与尺度在正式训练前冻结，不依据HR调参。')
    page();h('校准验收与图像域边界')
    p('固定合法train32按camera/frame字典序登记。校准结果k_R='+fmt(cal['k_R'],9)+'，lambda_G='+fmt(cal['lambda_G'],9)+'；G按真实LR有效位置梯度均方根（RMS）的0.1比例校准并在前500步渐入。R的校准不包含置信权重。系数来自原固定梯度统计，数值救援仅补影子更新验收，未重新搜索或计算系数。')
    rows=[]
    for method in ('LR_only','C1','R','G','RG'):
        observed=cost['shadow_delta_LR'][method]
        rows.append([method,*[sci(observed[scope]['LR_L1_delta']['mean']) for scope in ('self','same_time_view','same_view_time')]])
    table(['完整Adam单步','本观察LR L1变化','邻视角变化','邻时刻变化'],rows,[3.5,4.6,4.6,4.7])
    p('L1是平均绝对误差，MSE是均方误差。表中变化相对影子更新前父状态，负值表示L1降低。每种约束均独立复制完整模型、两套Adam状态、学习率和随机状态。R本观察LR输出变化RMS均值为'+sci(cost['shadow_delta_LR']['R']['self']['LR_output_change_RMS']['mean'])+'，并非零。因此图像梯度零空间不等于参数层面LR不变。该单步诊断只描述父状态附近，不是留出质量或长训练机制的证明。')
    p('完整画布CPU/3090降采样、伴随、Q幂等和反向泄漏检查通过；G稳定支持窗口的CUDA位置梯度、有限差分及协方差/不透明度停止梯度通过。最初样例含空射线，微扰使栅格支持改变；改为支持稳定窗口后验收，保留原失败、未放宽G登记容差、未修改训练损失。真实遮挡边缘仍可能不光滑。')
    p('校准保留'+str(len(data['calibration_failures']))+'次失败归档。后续完整Adam验收使用同路线重复计算的原生CUDA噪声基线和float32精度证据，联合/分开反向差异在记录范围内通过；通过记录不能删除先前失败及成本。图像监督和G梯度路由没有改变。这些检查属于校准阶段，该阶段的正式训练更新为0。核心正式训练另计72000步。')
    incident=data['source_incident']
    p('缓存启动后depth_prior.py曾追加来源/验收字段，实际加载初始源码与末尾文件哈希存在漂移。已重建实际源码并核对核心算法字节一致：初始SHA前12位'+incident['actual_loaded_source_sha256'][:12]+'，后续'+incident['later_file_sha256'][:12]+'。原收据保留，另以关联记录纠正身份，未静默改写历史。全部12端点的源码快照、父状态、两个Adam、随机状态、合法读取和检查点/浮点字节由final_integrity核验。')
    page();h('原因干预与历史资产恢复')
    cause=data['cause_interpretation'];counts=Counter(r['status'] for r in cause['cause_matrix_statuses'])
    p('三条各500步LR-only探针及冻结、颜色、配准、贡献支持诊断已完成。九项原因中：已确认实现问题'+str(counts['已确认实现问题'])+'，有干预支持的贡献因素'+str(counts['有干预支持的贡献因素'])+'，相关现象'+str(counts['相关现象'])+'，证据不足'+str(counts['证据不足'])+'。没有确认单一主因，不能用某项oracle收益宣布相机或几何错误。')
    probe=[]
    for method,step in [('U6000',0),('C1',500),('P_joint',500),('P_xyz',500),('P_SH',500)]:
        row=next(r for r in cause['probe_series'] if r['probe']==method and r['step']==step)
        probe.append([method,sci(row['train_lr_l1']),sci(row['train_lr_mse']),sci(row['cam01_hr_low_mse'])])
    table(['固定探针','train76 LR L1','train76 LR MSE','cam01 HR低频MSE'],probe,[3.1,4.6,4.6,5.1])
    p('P_joint相对C1的500步训练L1略降而MSE升高，不能确认SR妨碍真实LR优化。P_xyz/P_SH改善训练拟合，却使cam01低频误差升高；有限干预没有支持位置或直接球谐颜色是新视角误差的充分解释。探针训练于3090、C1于PRO6000，虽统一评价，微小差异仍受训练硬件影响。')
    p('P_xyz冻结形变参数，但改变输入位置仍可改变固定网络输出的颜色、尺度与opacity；因此这是参数位置限制，不是有效外观全部固定。P_SH仅改变直接base/children球谐系数，不能排除更广义的视角外观。冻结组参数和完整Adam状态按1/100/500步核验不变，children整体饱和比例很低，不能归因为普遍位置饱和。')
    p('固定2042个候选稳定支持点后，cam01的80/118帧低频MSE反而上升约1.251%/3.816%；关闭方向颜色与gain/bias也未支持简单修复。二维平移能减小部分残差，却可能补偿错误几何，仍只是目标开发LR诊断。cam01墙角3724贡献点有大量训练像素支持，cam02约覆盖95.028%目标贡献质量；这些贡献、厚度和大足迹均由父模型推断，不能当作唯一表面或米制几何真值。')
    p('P1已恢复固定dev8：复用4张登记浮点缓存，只补0/118的4张渲染，训练更新为0。它从头执行20200步、117971点；核心六臂从共同父状态续训、132972点。P1是历史质量参照，不是初始化、容量和优化轨迹配平的因果对照。原缺失报告与旧矩阵仍保留。上述开发LR/HR、oracle参数和诊断输出不进入训练权重或合法主成绩。')
    p('本地资产盘点显示Cook、Cut与MeetRoom的discussion/vrheadset准备缓存都只有0至118的60个偶数帧。Cook/Cut旧视频元数据为300帧，视频文件仍在，但本轮未重新完整解码。官方完整场景全集、完整基准与全时间先验尚未验收。Cut与MeetRoom旧先导已参与开发，不能直接作为此前未用于选择的独立确认场景。来源见future_benchmark_readiness.json。')
    page();h('实际成本与下一阶段边界')
    ledger=data['accounting'];core=ledger['actual']['core'];probes=ledger['actual']['cause_probe']
    p('完整账本记录核心正式更新'+str(core['completed_logged_updates'])+'，红绿蓝图像（RGB）前向'+str(core['logged_RGB'])+'，辅助矩前向'+str(core['logged_moments'])+'；原因探针更新'+str(probes['completed_logged_updates'])+'。每轮两个既有Adam分别服务相应参数组；这不意味着同一参数更新两次。缓存、检查、影子更新和评价成本另列，不能以相同迭代数宣称相同计算预算。')
    hardware=lambda name:'PRO6000' if 'PRO 6000' in name else 'A100' if 'A100' in name else '3090' if '3090' in name else name.replace('NVIDIA ','')
    resource=[]
    for method in stage.CORE:
        rr=[next(r for r in data['costs'] if r['endpoint']==f'r{repeat}_{method}') for repeat in ('1','2')]
        resource.append([method,*[hardware(r['training_gpu'])+' / '+fmt(r['training_seconds'],1) for r in rr],
                         fmt(rr[0]['peak_gb'],2)+' / '+fmt(rr[1]['peak_gb'],2)])
    table(['分支','后缀1硬件 / 训练秒','后缀2硬件 / 训练秒','峰值GB 1 / 2'],resource,[2.4,5.7,5.7,3.6])
    total=cost['total']
    pre=data['preprocessing'];recorded=pre['summary'];intervals=recorded['recorded_nonoverlapping_receipt_interval_seconds_by_resource_class']
    p('预处理账本含27阶段和76个哈希绑定引用。已记录且不重叠的CPU区间为'+fmt(intervals['CPU'],2)+'秒，GPU区间为'+fmt(intervals['GPU'],2)+'秒，CPU/GPU混合区间为'+fmt(intervals['CPU+GPU'],2)+'秒。合计'+fmt(recorded['recorded_interval_total_seconds_excluding_calibration_and_historical'],2)+'秒只是已记录区间之和，不是完整墙钟或并行完成时间。'+str(len(recorded['unknown_duration_stage_ids']))+'阶段未记录完整时长，相关峰值缺项也保留未知。历史资产复用不计本轮新增推理，不能视为免费。')
    p('子区间与外层控制器重叠区间不重复相加。校准失败、数值诊断和影子更新全部由calibration_total_cost.json单独负责：RGB '+str(total['rgb_forwards'])+'、矩 '+str(total['moment_forwards'])+'、影子轮次'+str(total['shadow_optimizer_rounds'])+'、Adam调用'+str(total['shadow_adam_calls'])+'、原生栅格反向'+str(total['native_rasterizer_backward_launches'])+'。校准阶段正式更新为'+str(total['formal_training_updates'])+'。救援部分记录'+fmt(cost['rescue_seconds'],2)+'秒与峰值'+fmt(cost['rescue_peak_gpu_bytes']/1e9,2)+'GB；校准完整墙钟未知。预处理账本只引用校准一次，不再累加其计数或时长。')
    if ledger.get('warnings'):p('账本仍显式保留警告：'+'；'.join(ledger['warnings']))
    decision=data['phase13']
    if decision:
        value=next((decision[k] for k in ('decision_zh','decision','conclusion_zh','conclusion','status') if k in decision),'已记录判断，见phase13_decision.json')
        p('第13节研究判断：'+(value if isinstance(value,str) else json.dumps(value,ensure_ascii=False))+'。该记录用于研究范围决定，不等于后续分支已完成。')
        reasons=next((decision[k] for k in ('reasons_zh','reason_zh','reasons','reason') if k in decision),None)
        if reasons:p('判断依据：'+('；'.join(map(str,reasons)) if isinstance(reasons,list) else str(reasons)))
        execution=decision.get('execution_status') or data.get('phase13_execution',{}).get('status')
        if execution:p('条件性扩展的实际执行状态：'+str(execution)+'。本报告已完成状态仅指核心六臂，不将进行中扩展写作终态。')
    else:p('第13节进入或停止判断尚未登记。核心完成不会自动启动动态或组件分支。只有存在有研究价值的质量收益/取舍且主要错误不是数据bug，才登记view/time/ST与mass对照及Full/Q二因素消融；稳定反证时保留负结果、收缩路线，不按任意0.1dB阈值或漂亮置信图决定。')
    p('本轮尚不支持跨场景一般性、米制几何正确或新时间泛化。方法与超参冻结后，需要至少两个此前未参与选择的完整场景，主方法与最强相同骨干基线各两次从头初始化。不同随机种子不能共享同一父模型后称独立训练。区域掩码只用于评价，保留背景、运动内部、显隐及失败案例。')
    p('完整固定图库保留cam00/01/02、40/80帧、两个后缀和所有登记ROI及六臂，见output/dynamic_sr_confidence_geometry_20261006/figures/index.json。下两页按预登记只展示frame40 suffix1的手部与墙角四格；没有按端点质量挑图。浮点指标与快速傅里叶变换（FFT）频带误差不受展示PNG量化影响。完整来源哈希、图库身份与两个成本账本保存在report_final_snapshot.json。')
    return blocks


def native_crops(data):
    """Preserve exact crop pixels; all labels are native Word paragraphs."""
    import numpy as np
    from PIL import Image
    out=RUN/'report_figures_final';out.mkdir(parents=True,exist_ok=True)
    result=[];roi=data['roi'];gallery=data['gallery']
    assert gallery['fixed_frames']==[40,80] and gallery['fixed_cameras']==['cam00','cam01','cam02']
    assert gallery['protocol_sha256']==stage.sha(RUN/'protocol.json') and roi['manifest_sha256']==data['protocol']['manifest']['sha256']
    for camera,region in FIGURE_SELECTION:
        x0,y0,x1,y1=roi['regions_by_camera_xyxy_exclusive'][camera][region]
        frame=40;paths=[DATA/'hr'/camera/'0040.png',*[RUN/'evaluation'/label/'extra/floats'/f'{camera}_0040.npz' for label in ('U6000','r1_C1','r1_RG')]]
        crops=[]
        for i,source in enumerate(paths):
            assert gallery['source_hashes'][str(source.relative_to(ROOT))]==stage.sha(source)
            if i==0:
                with Image.open(source) as img:array=np.asarray(img.convert('RGB'),dtype=np.uint8)
            else:
                meta=stage.read(source.with_suffix('.json'));assert stage.sha(source)==meta['sha256']
                cp=require_bound(data['protocol']['parent']) if i==1 else require_bound(stage.read(RUN/'runs'/('r1_C1' if i==2 else 'r1_RG')/'complete.json')['checkpoint'])
                assert meta['checkpoint_sha256']==stage.sha(cp)
                with np.load(source,allow_pickle=False) as saved:raw=saved['rgb_raw']
                assert raw.shape==(3,1008,1344) and np.isfinite(raw).all()
                array=np.rint(np.clip(raw.transpose(1,2,0),0,1)*255).astype(np.uint8)
            crop=array[y0:y1,x0:x1];assert crop.shape==(y1-y0,x1-x0,3)
            target=out/f'{camera}_0040_{region}_{i}.png';Image.fromarray(crop).save(target)
            crops.append(dict(label=PANEL_LABELS[i],path=str(target),sha256=stage.sha(target),source_path=str(source),
                              source_sha256=stage.sha(source),pixels=[x1-x0,y1-y0]))
        result.append(dict(camera=camera,frame=40,suffix='1',region=region,xyxy=[x0,y0,x1,y1],panels=crops,
            caption=('手与厨具局部细节' if camera=='cam00' else '灯与墙角低频失配')+'；同一登记坐标、原始裁剪像素，不锐化、不重采样文件。cam00/01均为已使用的开发视角。'))
    stage_save=RUN/'report_figures_final/index.json'
    stage_save.write_text(json.dumps(dict(status='prepared_pending_visual_review',selection='predeclared_before_core_completion',figures=result,
                                        gallery_source=data['source_identity']['gallery'],parameter_updates=0,gpu_forwards=0),ensure_ascii=False,indent=2)+'\n')
    return result


def append_figures(docx, figures):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm,Pt
    doc=Document(docx)
    for fig in figures:
        doc.add_page_break();doc.add_paragraph(('手与厨具' if fig['camera']=='cam00' else '灯与墙角')+'固定同帧对照','Heading 1')
        doc.add_paragraph(fig['camera']+'  frame40  suffix1  '+fig['region']+'  原裁剪'+str(fig['panels'][0]['pixels'][0])+'×'+str(fig['panels'][0]['pixels'][1]))
        doc.add_paragraph(fig['caption'])
        table=doc.add_table(rows=4,cols=2);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
        for col in table.columns:col.width=Cm(8.6)
        border=OxmlElement('w:tblBorders')
        for side in ('top','left','bottom','right','insideH','insideV'):
            element=OxmlElement('w:'+side);element.set(qn('w:val'),'single');element.set(qn('w:sz'),'4');element.set(qn('w:color'),'D9D9D9');border.append(element)
        table._tbl.tblPr.append(border)
        for i,panel in enumerate(fig['panels']):
            row=(i//2)*2;column=i%2;label=table.cell(row,column);photo=table.cell(row+1,column)
            label.text=panel['label'];label.paragraphs[0].runs[0].bold=True
            label.paragraphs[0].runs[0].font.size=Pt(11)
            label.paragraphs[0].paragraph_format.keep_with_next=True
            for cell in (label,photo):
                cell.width=Cm(8.6);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER;paragraph.paragraph_format.space_after=Pt(3)
                    paragraph.paragraph_format.space_before=Pt(3);paragraph.paragraph_format.line_spacing=Pt(15.5)
            # 168px hand crop is placed at native 96ppi. The wall retains all
            # 560x420 PNG pixels and is fitted to the cell without altering PNG.
            width=min(7.8,panel['pixels'][0]*2.54/96)
            photo.paragraphs[0].add_run().add_picture(panel['path'],width=Cm(width))
        for row in table.rows:
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        doc.add_paragraph('每幅显示文件保留原裁剪像素；Word按版面等比显示。参考HR只用于评价。四格分别为HR、共享父模型、同图完整SR、同图R与G；这两处是固定展示，不代表全图或另一后缀，完整图库与逐帧/区域结果共同约束结论。')
    doc.core_properties.title='动态场景超分置信度与软几何核心结果';doc.save(docx)


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--check-ready',action='store_true')
    parser.add_argument('--stem',type=Path,default=stage.DEFAULT_STEM);args=parser.parse_args()
    ready=readiness()
    if args.check_ready:print(json.dumps(ready,ensure_ascii=False));return
    if ready['status']!='ready_for_final_core_report':raise RuntimeError(ready)
    data=snapshot()
    for item in data['finalization'].values():
        if isinstance(item,dict) and 'path' in item and 'sha256' in item:require_bound(item)
    for item in data['cause_interpretation']['asset_integrity']:require_bound(item)
    p1=data['P1_recovery'];assert p1['parameter_updates']==0
    for key in ('plan','float_index','cached_report','matrix'):require_bound(p1[key])
    preprocessing_check=verify_preprocessing(data['preprocessing'])
    assert data['calibration_total']['status']=='passed_all_calibration_checks'
    assert data['calibration_total']['coefficients_recalculated_on_rescue'] is False
    blocks=content(data);figures=native_crops(data);args.stem.parent.mkdir(parents=True,exist_ok=True)
    stage.save_markdown(args.stem.with_suffix('.md'),blocks);stage.save_docx(args.stem.with_suffix('.docx'),blocks)
    append_figures(args.stem.with_suffix('.docx'),figures)
    with args.stem.with_suffix('.md').open('a') as stream:
        for figure in figures:
            stream.write('\n## '+figure['camera']+'固定同帧对照\n\n'+figure['caption']+'\n\n')
            for panel in figure['panels']:stream.write(panel['label']+'\n\n!['+panel['label']+']('+panel['path']+')\n\n')
        stream.write('[完整固定图库]('+str(RUN/'figures/index.json')+')\n')
    receipt=dict(status='core_completed_report_pending_each_page_visual_review',created_at_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),
        source_sha256=stage.sha(__file__),source_identity=data['source_identity'],readiness=ready,preprocessing_validation=preprocessing_check,
        docx_path=str(args.stem.with_suffix('.docx')),docx_sha256=stage.sha(args.stem.with_suffix('.docx')),
        md_path=str(args.stem.with_suffix('.md')),md_sha256=stage.sha(args.stem.with_suffix('.md')),
        figure_index=dict(path=str(RUN/'report_figures_final/index.json'),sha256=stage.sha(RUN/'report_figures_final/index.json')),
        visual_QA='pending_render_and_each_page_inspection',scope='Completed core development experiment; independent-scene/phase13 evidence is not implied')
    (RUN/'report_final_snapshot.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(status=receipt['status'],docx=receipt['docx_path']),ensure_ascii=False))


if __name__=='__main__':main()
