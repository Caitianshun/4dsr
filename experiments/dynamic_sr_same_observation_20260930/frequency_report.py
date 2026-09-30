"""Append the mml-compatible FFT audit and bounded constraint interpretation."""
import json
from pathlib import Path
import shutil
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output/dynamic_sr_same_observation_20260930/frequency'
DOC = ROOT / 'docs/dynamic_sr_prior_diagnosis_results_2026-09-29.docx'
MARKER = '参照mml的FFT频带误差补充'
BANDS = ['low', 'mid', 'high']
ARMS = ['LR-direct-HRrender', 'LR-direct-Bicubic', 'HR-direct-6k', 'U6000', 'J1', 'Async2', 'Sync2', 'T', 'P1-from-start']


def read(p): return json.loads(Path(p).read_text())
def fmt(x, n=3): return f'{x:.{n}f}'
def sub(base, label):
    node = OxmlElement('m:sSub'); e = OxmlElement('m:e'); s = OxmlElement('m:sub')
    if isinstance(base, str): e.append(mrun(base))
    else:
        for b in base: e.append(mrun(b) if isinstance(b, str) else b)
    s.append(mrun(label)); node.extend([e, s]); return node
def sup(base, power):
    node = OxmlElement('m:sSup'); e = OxmlElement('m:e'); s = OxmlElement('m:sup')
    if isinstance(base, str): e.append(mrun(base))
    else:
        for b in base: e.append(mrun(b) if isinstance(b, str) else b)
    s.append(mrun(power)); node.extend([e, s]); return node
def fraction(numerator, denominator):
    node = OxmlElement('m:f'); n = OxmlElement('m:num'); d = OxmlElement('m:den')
    for x in numerator: n.append(mrun(x) if isinstance(x, str) else x)
    for x in denominator: d.append(mrun(x) if isinstance(x, str) else x)
    node.extend([n, d]); return node
def mrun(text):
    r = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = text; r.append(t); return r


def append_report(doc, md):
    s = read(OUT / 'summary.json'); check = read(OUT / 'operator_check.json')
    assert s['status'] == 'completed_frequency_audit'
    assert not any(p.text == MARKER for p in doc.paragraphs), 'Frequency supplement already present'
    def text(value, small=False):
        p = doc.add_paragraph(value)
        if small:
            for run in p.runs: run.font.size = Pt(9)
        md.extend([value, '']); return p
    def heading(value):
        p = doc.add_heading(value, 1); p.paragraph_format.page_break_before = True
        md.extend(['<!-- pagebreak -->', '', '## ' + value, ''])
    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False
        for col, width in zip(t.columns, widths): col.width = Cm(width)
        for cell, value in zip(t.rows[0].cells, headers): cell.text = str(value)
        for row in rows:
            for cell, value in zip(t.add_row().cells, row): cell.text = str(value)
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
        t._tbl.tblPr.append(borders)
        for ri, row in enumerate(t.rows):
            prop = row._tr.get_or_add_trPr(); prop.append(OxmlElement('w:cantSplit'))
            if ri == 0: prop.append(OxmlElement('w:tblHeader'))
            for ci, cell in enumerate(row.cells):
                cell.width = Cm(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = cell._tc.get_or_add_tcPr(); shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '233B53' if ri == 0 else 'EEF3F7' if ri % 2 else 'FFFFFF'); pr.append(shade)
                margin = OxmlElement('w:tcMar')
                for side, val in [('top', '45'), ('bottom', '45'), ('left', '65'), ('right', '65')]:
                    e = OxmlElement('w:' + side); e.set(qn('w:w'), val); e.set(qn('w:type'), 'dxa'); margin.append(e)
                pr.append(margin)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs: run.font.size = Pt(9.5); run.bold = ri == 0; run.font.color.rgb = RGBColor(255, 255, 255) if ri == 0 else RGBColor(0, 0, 0)
        p = doc.add_paragraph(); p.paragraph_format.line_spacing = Pt(3); p.paragraph_format.space_after = Pt(4)
        md.extend(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |'] + ['| ' + ' | '.join(map(str, row)) + ' |' for row in rows] + [''])
    def equation(nodes, tex):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para = OxmlElement('m:oMathPara'); math = OxmlElement('m:oMath'); para.append(math)
        for n in nodes: math.append(mrun(n) if isinstance(n, str) else n)
        p._p.append(para); md.extend(['$$', tex, '$$', ''])
    def picture(path, caption):
        doc.add_picture(str(path), width=Cm(17.5)); doc.paragraphs[-1].paragraph_format.keep_with_next = True
        p = doc.add_paragraph(caption, 'Caption'); p.paragraph_format.keep_with_next = False
        md.extend([f'![{caption}]({path})', ''])
    def model(arm, scope='equal_camera_mean'):
        return next(r for r in s['summaries'] if r['arm'] == arm and r['scope'] == scope)
    def gap(arm, scope='equal_camera_mean'):
        return next(r for r in s['p1_gaps'] if r['reference'] == arm and r['scope'] == scope)

    heading(MARKER)
    text('2026年9月30日补充。按照用户提供的mml日志及本机对应源码，对全部九种模型或输出新增径向FFT误差分解，包含P1从头训练和历史两套后缀。结果支持优先诊断cam01低频重建误差；cam00还需检查中高频细节。频带定位不能单独证明几何原因，以下训练约束均区分已使用与待验证。')
    text('参考文件：/home/cai_tianshun/Project/mml/scripts/freq_band_decomp.py，第36至43行定义径向环带，第52至63行完成误差变换与归一化，第64至74行汇总。说明日志为docs/week3_freq_band_diagnostic.md；hf_lrgate_loss_walkthrough.md只参考通用约束思路，不迁入其人体模板、轮廓、抠背景或INTER_AREA退化假设。三个来源的SHA256已登记。', small=True)
    text('输入是同一相机、同一时间的HR尺寸预测R与真实HR图Y。先计算RGB残差e=R−Y，再做二维快速傅里叶变换FFT；它把变化缓慢的误差与细纹理误差分开。这里分析残差，不分析图像内容的占比，也没有把原始LR图片直接混入HR残差计算。')
    equation([sub('E', 'b'), ' = ', fraction([sup([sub(['‖', sub('M', 'b'), ' ⊙ FFT2(R − Y)‖'], '2')], '2')], ['3', sup('(HW)', '2')])],
             r'E_b=\frac{\lVert M_b\odot\operatorname{FFT2}(R-Y)\rVert_2^2}{3(HW)^2}')
    text('H、W是HR图高和宽，3是RGB通道数。M是对应频带的0/1掩码，⊙表示逐系数相乘；分子对所有频率与通道求平方和。NumPy默认FFT前向不归一化，因此分母是3(HW)²。三个频带互不重叠且覆盖全部系数，其E之和等于全图MSE。零频率DC包含在low，不去均值、不加窗。')
    table(['频带','径向频率r','读法'], [['low','0 ≤ r < 0.125','缓慢结构、亮度及色偏误差'], ['mid','0.125 ≤ r < 0.25','中等尺度边缘与纹理误差'], ['high','r ≥ 0.25','细边缘与纹理误差']], [2.2, 4.6, 10.7])
    text('r=√(fx²+fy²)，单位为周期/HR像素；0.125=0.5/4是×4的名义单轴Nyquist频率。二维LR采样的名义频率范围是方形，而这里是圆环；真实bicubic还包含滤波、边界、clamp和数据量化，因此不能将这三圈直接当作真实LR算子的可观测空间与零空间。')
    text('4dsr原分析是正交DCT的方形频带，单独分离DC；本补充是FFT径向环带并将DC计入low，二者数值不能混为同一列。mml读取渲染PNG；本项目保留主评价的clamp后float32渲染，HR仍从PNG/255读取，不增加预测量化。相同PNG测试图已逐项复现mml算术。', small=True)

    heading('全模型FFT频带结果')
    text('cam00和cam01各60个登记帧，J1、Async2、Sync2和T均纳入两套续训，其余为单个端点。共13端点、1560对图像。总PSNR与各频带PSNR均先逐帧计算再平均；占比则先累计MSE再求比例。频带PSNR的峰值仍取1，分母仍按整图归一化，因此三个频带PSNR不能相加。')
    table(['模型或输出','总PSNR','low PSNR','mid PSNR','high PSNR','L / M / H %'],
          [[a, fmt(model(a)['psnr']), *[fmt(model(a)[b + '_psnr']) for b in BANDS], ' / '.join(fmt(model(a)[b + '_share_pct'], 2) for b in BANDS)] for a in ARMS], [4.2, 2.1, 2.1, 2.1, 2.1, 4.9])
    equation([sub('PSNR', 'n,b'), ' = −10 log₁₀(max(', sub('E', 'n,b'), ', ', sup('10', '−12'), '))'],
             r'\operatorname{PSNR}_{n,b}=-10\log_{10}\bigl(\max(E_{n,b},10^{-12})\bigr)')
    equation([sub('share', 'b'), ' = ', fraction([sub('Σ', 'n'), ' ', sub('E', 'n,b')], [sub('Σ', 'n'), ' ', sub('Σ', 'k'), ' ', sub('E', 'n,k')]), ' × 100%'],
             r'\operatorname{share}_b=\frac{\sum_n E_{n,b}}{\sum_n\sum_k E_{n,k}}\times100\%')
    text('n表示一对登记图像，b、k表示频带。表内PSNR对n算术平均；占比先汇总每帧误差能量，因此不会把一个几乎无误差的帧与一个大误差帧的百分比同等对待。归一化使用相同HR尺寸和RGB三通道，不除以频带内系数数量。')
    heading('绝对频带误差与占比的关系')
    text('下表为绝对频带MSE贡献，所有数值均乘1000显示。例如P1的low数值3.08331代表MSE=0.00308331。占比变化必须结合绝对误差判断；模型越模糊，也可能有较低的高频误差，不能据此认定其细节更真实。', small=True)
    table(['模型或输出','low ×1000','mid ×1000','high ×1000'],
          [[a, *[fmt(model(a)[b + '_mse'] * 1000, 5) for b in BANDS]] for a in ARMS], [6.1, 3.8, 3.8, 3.8])
    j = model('J1'); new = model('P1-from-start')
    text('J1的低频占比为'+fmt(j['low_share_pct'], 2)+'%，P1为'+fmt(new['low_share_pct'], 2)+'%。P1三个频带绝对MSE均增加，分别比J1高'+ '、'.join(fmt(gap('J1')[b + '_mse_relative_pct'], 2)+'%' for b in BANDS)+'。P1高频占比从'+fmt(j['high_share_pct'], 2)+'%降至'+fmt(new['high_share_pct'], 2)+'%，同时高频绝对误差增加；这正是仅看比例会误判的例子。')
    text('两相机帧数相等，因此总MSE为相机等权；能量占比仍受高误差相机主导，并非两个相机百分比的算术平均。PSNR均值与“平均MSE再转PSNR”不同，机器记录保留两者。主表PSNR复现最大差为'+f"{s['max_main_psnr_gap']:.2e}"+' dB。', small=True)

    heading('分相机定位P1新增误差')
    picture(OUT / 'band_mse_by_camera.png', '图中堆叠的是绝对MSE贡献。两相机横轴刻度不同，不能按条形长度直接跨面板比较。')
    table(['P1减参照','low MSE差 ×1000','mid差 ×1000','high差 ×1000','low占新增误差 %'],
          [[a, *[fmt(gap(a)[b + '_mse_delta'] * 1000, 6) for b in BANDS], fmt(gap(a)['low_share_of_total_delta_pct'], 2)] for a in ['U6000', 'J1', 'Async2', 'Sync2', 'T']], [3.3, 3.8, 3.1, 3.1, 4.2])
    a, b = gap('J1', 'cam00'), gap('J1', 'cam01')
    text('P1相对J1：cam00新增MSE='+f"{a['mse_delta']:.6e}"+'，低/中/高频贡献为'+ ' / '.join(fmt(a[x + '_share_of_total_delta_pct'], 2) for x in BANDS)+'%；cam01新增MSE='+f"{b['mse_delta']:.6e}"+'，对应为'+ ' / '.join(fmt(b[x + '_share_of_total_delta_pct'], 2) for x in BANDS)+'%。两相机合并后，低频贡献'+fmt(gap('J1')['low_share_of_total_delta_pct'], 2)+'%的新增误差。')
    text('直观上，cam01要先处理大尺度位置、亮度或外观偏差；cam00的变化更多涉及边缘与细纹理。频率只能说明误差变化的尺度：遮挡/几何错位、曝光、颜色和优化不足都可能产生低频误差，锐利边缘的几何错位也能产生高频误差，不能由此决定冻结几何或更换运动模型。')
    text('现有PSNR、SSIM、LPIPS与固定墙角/衣纹图像结论继续保留。本补充没有重新挑选展示帧，没有训练更新，也不把两相机开发样例升格为独立场景验证。', small=True)

    heading('教师误差与真实LR约束')
    text('在四个既有教师训练相机cam02、cam06、cam12、cam18及第0、40、80、118帧上，补充16观察。HR只用于诊断教师质量；这些数值与前页开发相机结果分开，不混用旧train76均值。LR-bicubic此处是“真实LR图上采样”，与主表的“LR模型在LR画布渲染后上采样”不同。')
    table(['同位置图像对','总PSNR','low MSE ×1000','mid ×1000','high ×1000'],
          [[r['arm'], fmt(r['psnr']), *[fmt(r[b + '_mse'] * 1000, 6) for b in BANDS]] for r in s['train16']], [4.8, 2.3, 3.6, 3.4, 3.4])
    table(['降采样后与真实LR比较','RGB平均L1','MSE'],
          [[r['arm'], fmt(r['lr_l1'], 7), f"{r['lr_mse']:.7e}"] for r in s['train16_closure']], [6.5, 5.5, 5.5])
    text('SwinIR对HR的低频MSE为4.9887×10⁻⁶，而P1为1.8653×10⁻⁴；真实LR闭环L1分别为0.0005786和0.0074168。固定训练位置上，教师低频已较准确，模型仍存在更大的观测拟合误差。这支持检查重建和优化环节，不支持直接宣布“教师低频偏差是P1退步的主因”。')
    text('进一步核验H(T)=T−Up(D₀(T))。D₀为本项目未clamp、未量化的线性bicubic降采样，Up也采用既定bicubic。16张真实SwinIR目标的D₀(H(T))平均绝对值='+fmt(check['mean_d0_resize_residual_l1'], 7)+'，MSE='+f"{check['mean_d0_resize_residual_mse']:.7e}"+'，并非0。“降采样核相同”本身不能保证D₀Up是恒等映射，因此该残差不是严格零空间投影。')
    text('一个明确的方向性例子：fx=fy=0.109375的斜纹位于LR名义Nyquist方形内，却被径向FFT划入mid；经真实D₀后RMS仍为0.29316，HR输入RMS为0.70711。mid里也存在LR能响应的成分，所以“仅监督mid/high”仍可能改变LR输出。', small=True)
    text('旧76观察的8步真实伴随回投影已将闭环MSE从6.3876×10⁻⁷降到9.5943×10⁻¹⁰，但PSNR变化−0.00417 dB、SSIM−0.000227、LPIPS−0.003535。更好闭环没有自动带来更好HR像素结果；不能因为本次低频占比高而忽略这项已有证据。', small=True)

    heading('这些信息能形成什么训练约束')
    text('最直接的约束仍是真实LR观测一致性：同一时刻、同一相机渲染一张HR图，再通过匹配数据的降采样算子与真实LR比较。它约束模型必须解释实际观测，已在P1中以权重1使用。本次不是发现一个缺失的LR项，而是发现“已有LR项仍未保证充分拟合”，需要先检查优化和与SR项的相互作用。')
    equation([sub('L', 'LR'), ' = mean|D(', sub('R', 'θ'), ') − y|'],
             r'\mathcal{L}_{\mathrm{LR}}=\operatorname{mean}\lvert D(R_\theta)-y\rvert')
    text('Rθ是参数θ渲染的HR图，y是对应真实训练LR；D沿用bicubic抗混叠及clamp，预测不量化取整。真实HR和开发相机图像不能进入这项训练目标。')
    text('可做的最小候选是将SR辅助损失限定在中高频残差，而保留完整LR项：让教师主要提供纹理建议，减少它对缓慢颜色与结构的直接牵引。以下Pdetail是FFT逆变换后的mid+high残差，阈值暂按本次诊断协议；这是候选方法，不是已训练的结果。')
    equation([sub('L', 'candidate'), ' = ', sub('L', 'LR'), ' + ', sub('λ', 'SR'), ' mean|', sub('P', 'detail'), '(', sub('R', 'θ'), ' − T)| + ', sub('L', 'reg')],
             r'\mathcal{L}_{\mathrm{candidate}}=\mathcal{L}_{\mathrm{LR}}+\lambda_{\mathrm{SR}}\operatorname{mean}\lvert P_{\mathrm{detail}}(R_\theta-T)\rvert+\mathcal{L}_{\mathrm{reg}}')
    text('T是相同训练LR生成的冻结SwinIR目标；Pdetail(X)=IFFT2[(1−Mlow)⊙FFT2(X)]。两项仍可由一次HR渲染、一次联合反传、一次更新实现；形变、几何和外观参数默认均可收到梯度。频带滤波增加计算，但不需要第二次渲染。该投影并不保证零LR干扰，也可能削弱教师的有益中尺度信息或引入周期边界振铃，因此必须由受控验证判断。')
    text('λSR不能由误差占比直接推导。P1原全图SR权重0.1、Async2/Sync2两项各0.05仍是既有设置；新增滤波会改变残差尺度与参数梯度，即使沿用0.1也不等于同一梯度强度。第一步应报告各项损失及对运动/位置/外观的梯度大小，不能将95.74%低频新增误差换成95.74%的LR权重。')
    text('若确实要保持LR观测不受目标修正影响，应依据真实线性D₀及其伴随构造投影，而非凭FFT阈值或Up替代伴随。下式为理想目标锚定：†表示Moore–Penrose伪逆，不要求DDᵀ可逆。')
    equation([sub('T', 'anchor'), ' = T + ', sup('D₀', 'T'), sup(['(D₀', sup('D₀', 'T'), ')'], '†'), '(y − D₀T)'],
             r'T_{\mathrm{anchor}}=T+D_0^{\mathsf T}(D_0D_0^{\mathsf T})^{\dagger}(y-D_0T)')
    text('y若在D₀的值域中，精确求解后D₀Tanchor=y。有限迭代只是近似，clamp会破坏线性等式；目标满足闭环也不保证更新后的模型满足闭环或高频真实。旧8步锚定已有取舍结果，本补充不宣布重新采用它。', small=True)

    heading('约束的适用边界与最小验证')
    table(['候选约束','合法依据与作用','尚需验证'],
          [['真实LR闭环','已有训练LR和既定退化算子；约束可观测颜色结构','P1已有此项，检查优化与参数梯度冲突'],
           ['SR频带选择','训练LR生成的教师；减少其低频直接监督','中高频拟合、LR闭环、振铃及新视角性能'],
           ['教师置信门控','训练LR闭环偏差或可见性可靠的多视角对应','低LR偏差不能证明高频真实；遮挡须排除'],
           ['SR梯度路由','训练观测上的运动/外观梯度及跨时间响应','频带能量不足以决定冻结几何或共享运动'],
           ['时序和跨视角一致性','合法训练视角、时间、标定与可见性检查','错误细节也可能一致；需独立真实质量检验']], [3.3, 7.1, 7.1])
    text('可执行的下一项最小验证：从同一U6000检查点及完整优化器/RNG/采样状态出发，固定点数、相机时间序列、教师、损失外部系数和6000步后缀，仅将J1全图SR残差改为固定mid+high滤波残差。保持既有LR权重1，记录滤波后的SR梯度尺度。它检验频带选择的作用，不检验“从第一步用SR”的效果；若随后检验起始时机，应另做配对从头训练并控制增密和SR曝光。此次未启动该训练。')
    text('支持该候选需要：真实训练LR闭环改善，cam01低频绝对MSE下降，同时cam00细节、SSIM、LPIPS和时序没有不可接受的损失；所有取舍如实报告。若只改善闭环而HR质量不改善，或损害cam00，不把候选升级为主线，也不为解释结果追加门控网格。几何/运动因果判断还需独立梯度路由消融，频谱表不能代替。')
    text('cam00、cam01均为已使用的开发证据；用本次统计指导方法后，最终有效性须在独立固定完整场景确认。HR误差图、HR光流、频带占比都不能直接生成LR-only训练掩码。新的阈值、门控或权重若依赖HR开发成绩选择，必须披露开发信息，不能把它写成真实LR自行推导的约束。')
    text('复算完整性：复用600对已验SHA256的浮点缓存，新做976次RGB渲染（含16训练位置），新增模型更新0；两独立作业墙钟约94.7与151.6秒并行执行，另计CPU算子检查及图表/报告。最终复算使用本机RTX3090，旧缓存硬件沿原receipt保留；不把诊断前向计入训练预算。Parseval最大误差='+f"{s['max_parseval_gap']:.2e}"+'。', small=True)
    text('数值与来源：output/dynamic_sr_same_observation_20260930/frequency下的protocol.json、verification.json、endpoints/*.json、per_frame.csv、model_summary.csv、p1_band_gaps.csv、train16.json、operator_check.json、summary.json和figures.json。旧DCT报告及训练记录未改数值；本次文件哈希与Word视觉核验单独存档，不以旧23页报告的哈希表示更新后的文档。', small=True)
    text('算法依据：NumPy FFT官方文档 https://numpy.org/doc/stable/reference/routines.fft.html 说明变换归一化、幅度/相位与功率谱；实际降采样以本项目冻结PyTorch实现及数值伴随检查为准。', small=True)


def main():
    backup = OUT / 'document_before_update'
    if backup.exists():
        for p in [DOC, DOC.with_suffix('.md'), DOC.with_suffix('.codex.md')]:
            assert p.read_bytes() == (backup / p.name).read_bytes(), 'Only retry an unchanged document after authoring failure'
    else:
        backup.mkdir()
        for p in [DOC, DOC.with_suffix('.md'), DOC.with_suffix('.codex.md')]: shutil.copy2(p, backup / p.name)
    doc = Document(DOC); md = DOC.with_suffix('.md').read_text().splitlines()
    append_report(doc, md); doc.save(DOC); DOC.with_suffix('.md').write_text('\n'.join(md) + '\n')
    print(DOC)


if __name__ == '__main__': main()
