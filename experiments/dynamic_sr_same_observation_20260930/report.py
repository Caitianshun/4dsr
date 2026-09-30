"""Append the completed from-start comparison to the existing validation document.

Run with the bundled artifact Python runtime after the DOCX edit operation marker.
The source Markdown and existing report generator share this append function.
"""
from copy import deepcopy
import csv
import json
from pathlib import Path
import shutil
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt,RGBColor

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'output/dynamic_sr_same_observation_20260930'
DOC=ROOT/'docs/dynamic_sr_prior_diagnosis_results_2026-09-29.docx'
MARKER='从初始化开始的同图联合监督验证'
ARMS=['LR-direct-HRrender','LR-direct-Bicubic','HR-direct-6k','U6000','J1','Async2','Sync2','T','P1-from-start']


def read(path):return json.loads(Path(path).read_text())
def fmt(v,n=6):return f'{float(v):.{n}f}'
def delta(v,n=6):return f'{float(v):+.{n}f}'


def append_report(doc,md):
    s=read(OUT/'comparison_summary.json');assert s['status']=='completed_comparison'
    assert not any(p.text==MARKER for p in doc.paragraphs),'This completed comparison is already present'
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                if '2026-09-29' in run.text:run.text=run.text.replace('2026-09-29','2026-09-30')
    new=next(r for r in s['new_run'] if r['scope']=='equal_camera_mean')
    historical=read(ROOT/'output/dynamic_sr_prior_diagnosis_20260929/final_summary.json')['averages']
    averages=[next(r for r in historical if r['arm']==a and r['scope']=='equal_camera_mean') if a!='P1-from-start' else new for a in ARMS]
    # Preserve the original first table and row formatting, adding one completed model.
    first=doc.tables[0]
    assert first.rows[0].cells[0].text=='模型或输出'
    clone=deepcopy(first.rows[-1]._tr);first._tbl.append(clone)
    for cell,value in zip(first.rows[-1].cells,['P1-from-start',fmt(new['psnr'],4),fmt(new['ssim']),fmt(new['lpips'])]):
        pr=cell._tc.get_or_add_tcPr()
        for old_shade in list(pr.findall(qn('w:shd'))):pr.remove(old_shade)
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEF3F7');pr.append(shade)
        paragraph=cell.paragraphs[0]
        if paragraph.runs:
            paragraph.runs[0].text=value
            for run in paragraph.runs[1:]:run.text=''
        else:paragraph.add_run(value)
    notice=doc.add_paragraph('P1-from-start为9月30日追加的独立从头训练，共20200步，第一步起同图LR与SR监督；仅一次训练，详见文末追加验证。历史两套续训的说明不适用于P1。')
    for run in notice.runs:run.font.size=Pt(9)
    first._tbl.addnext(notice._p)
    start=next(i for i,line in enumerate(md) if line.startswith('| 模型或输出 |'))
    stop=start+1
    while stop<len(md) and md[stop].startswith('|'):stop+=1
    md.insert(stop,'| P1-from-start | '+fmt(new['psnr'],4)+' | '+fmt(new['ssim'])+' | '+fmt(new['lpips'])+' |')
    md.insert(stop+2,'2026年9月30日追加P1-from-start：从原始LR点云初始化，第一步起同图LR与SR联合监督；单次从头训练，与历史两套续训均值分开解释。详细设置和成本见文末追加验证。')
    md.insert(stop+3,'')

    def text(value,small=False):
        p=doc.add_paragraph(value)
        if small:
            for r in p.runs:r.font.size=Pt(9)
        md.extend([value,'']);return p
    def heading(value):
        p=doc.add_heading(value,1);p.paragraph_format.page_break_before=True
        md.extend(['<!-- pagebreak -->','','## '+value,''])
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for c,w in zip(t.columns,widths):c.width=Cm(w)
        for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
        for values in rows:
            for c,v in zip(t.add_row().cells,values):c.text=str(v)
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        t._tbl.tblPr.append(borders)
        for ri,row in enumerate(t.rows):
            prop=row._tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
            if ri==0:prop.append(OxmlElement('w:tblHeader'))
            for ci,cell in enumerate(row.cells):
                cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'233B53' if ri==0 else 'EEF3F7' if ri%2 else 'FFFFFF');pr.append(shade)
                margin=OxmlElement('w:tcMar')
                for side,val in [('top','30'),('bottom','30'),('left','75'),('right','75')]:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margin.append(e)
                pr.append(margin)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:r.font.size=Pt(9.5);r.bold=ri==0;r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
        spacer=doc.add_paragraph()
        spacer.paragraph_format.line_spacing=Pt(3)
        spacer.paragraph_format.space_before=Pt(0)
        spacer.paragraph_format.space_after=Pt(4)
        md.extend(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+
                  ['| '+' | '.join(str(v) for v in row)+' |' for row in rows]+[''])
    def picture(path,caption):
        doc.add_picture(str(path),width=Cm(17.5));doc.paragraphs[-1].paragraph_format.keep_with_next=True
        p=doc.add_paragraph(caption,'Caption');p.paragraph_format.keep_with_next=False
        md.extend([f'![{caption}]({path})',''])
    def equation():
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para=OxmlElement('m:oMathPara');math=OxmlElement('m:oMath');para.append(math)
        def run(parent,value):
            r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;r.append(t);parent.append(r)
        def term(label):
            node=OxmlElement('m:sSub');base=OxmlElement('m:e');sub=OxmlElement('m:sub')
            run(base,'L');run(sub,label);node.append(base);node.append(sub);math.append(node)
        term('total');run(math,' = ');term('LR');run(math,' + 0.1 ');term('SR');run(math,' + ');term('reg')
        p._p.append(para)
        md.extend(['$$',r'\mathcal{L}_{\mathrm{total}}=\mathcal{L}_{\mathrm{LR}}+0.1\mathcal{L}_{\mathrm{SR}}+\mathcal{L}_{\mathrm{reg}}','$$',''])

    heading(MARKER)
    text('2026年9月30日追加验证。P1-from-start按用户提出的方式执行一次独立训练：从合法训练LR构造的20712点原始点云及新模型参数出发，每步只渲染一张HR图，在同一相机、同一时间上同时计算真实LR与冻结SwinIR目标的误差，合并损失后一次反向传播和一次Adam更新。没有加载U6000或其他已训练的重建检查点。')
    equation()
    text('LR项是同一HR渲染经过既有bicubic抗混叠降采样与clamp后，对真实LR的全图RGB平均绝对误差；SR项是该HR渲染对对应SwinIR PNG目标的全图RGB平均绝对误差。正则项沿用细训练阶段的时间平滑、时间平面L1和空间平面总变分。降采样过程保留梯度，不加入PNG量化取整。')
    text('细训练正则系数依次为0.001、0.0001、0.0002；粗训练不计算这三项。图像损失没有加入SSIM或LPIPS项。这两个指标用于评价。1.0与0.1是损失缩放系数，不是实际梯度贡献的比例：目标误差和图像到模型参数的导数都会影响各参数的更新。',small=True)
    table(['训练阶段','更新数','LR权重','SR权重','每步渲染 反传 更新'],
          [['粗训练','1000','1.0','0.1','1 / 1 / 1'],['动态细训练','19200','1.0','0.1','1 / 1 / 1']],
          [3.7,2.3,2.3,2.3,6.9])
    text('粗训练沿用骨干的静态高斯初始化策略，形变网络在细训练阶段启用；两个阶段从各自第一步都同时使用LR与SR。采样覆盖cam02至cam20共19个训练相机、60个登记帧，每步均匀抽取一个相机与帧的组合。训练输入仅为真实训练LR及其固定SwinIR目标；真实HR和cam00、cam01只用于评价。')
    text('总预算在看结果前固定为20200次更新，以对齐J1、Async2、Sync2、T的累计更新数。沿用原粗细训练的优化器重置、学习率与增密规则；点数阈值120000为软上限，实际终点为'+str(s['training']['points'])+'点。与132972点的旧分裂表示不同，因此这是整套训练流程对照，而不是只改变反传次数的单因素消融。')
    text('粗训练和细训练的首步审计均确认：同一渲染张量用于两项监督，两者输出梯度非零，联合梯度等于LR梯度加0.1倍SR梯度。训练循环静态核查和实际计数均为每步一次渲染、一次参数反传、一次更新。7000、14200、20200步检查点都已回传并核对SHA256，20200为唯一预定主终点。',small=True)

    heading('同图联合监督与现有模型的指标比较')
    text('主成绩在同一本机RTX3090上，复用原始浮点渲染评价、相同clamp、PSNR、SSIM及AlexNet v0.1 LPIPS算术。cam00、cam01各60帧先逐帧平均，再对两相机等权；新方法为一次从头训练，J1、Async2、Sync2、T为共享U6000的两套续训均值。两相机已参与研究开发。')
    table(['模型或输出','PSNR dB','SSIM','LPIPS'],
          [[r['arm'],fmt(r['psnr'],4),fmt(r['ssim']),fmt(r['lpips'])] for r in averages],[6.5,3.5,3.5,4.0])
    gaps=[r for r in s['comparisons'] if r['scope']=='equal_camera_mean' and r['reference'] in ['U6000','J1','Async2','Sync2','T']]
    table(['P1减参照','PSNR差 dB','SSIM差','LPIPS差'],
          [[r['reference'],delta(r['psnr'],4),delta(r['ssim']),delta(r['lpips'])] for r in gaps],[6.5,3.5,3.5,4.0])
    j=next(r for r in gaps if r['reference']=='J1')
    text('本次P1相对J1的PSNR变化为'+delta(j['psnr'],4)+' dB，SSIM变化为'+delta(j['ssim'])+'，LPIPS变化为'+delta(j['lpips'])+'；相对U6000、J1、Async2、Sync2和T，三项主均值均变差。此次实验不支持以P1替换现有分阶段方案。',small=True)
    text('PSNR和SSIM的正差有利，LPIPS的负差有利。三指标分别判断，不合成为单一分数。该次结果同时改变了监督起始时机、LR与SR的采样耦合、增密轨迹和SR曝光次数；不把改进或退步单独归因于同图联合损失，也不由一次训练建立重复性或跨场景结论。',small=True)
    text('权重对照：U6000和J1均为LR 1.0、SR 0.1；Async2和Sync2为LR 1.0、两张SR各0.05。T的LR仍为1.0，SR外部系数0.1，但像素误差另乘固定需求图及8.3434504的初始梯度校准倍率，所以T不等同于全图均匀0.1。P1的SR没有这项像素选择或校准。',small=True)

    heading('分相机结果与训练预算')
    table(['模型与相机','PSNR dB','SSIM','LPIPS'],
          [[r['arm']+' / '+r['scope'],fmt(r['psnr'],4),fmt(r['ssim']),fmt(r['lpips'])]
           for a in ['J1','Async2','Sync2','T','P1-from-start']
           for r in (s['new_run'] if a=='P1-from-start' else historical) if r['arm']==a and r['scope'] in ['cam00','cam01']],
          [6.5,3.5,3.5,4.0])
    table(['模型','累计更新','RGB渲染数','SR图像数','点数'],
          [[r['arm'],r['total_updates'],r['training_rgb_forwards'],r['SR_observations'],r['points']] for r in s['training_budgets']],
          [5.7,2.8,3.0,3.0,3.0])
    text('预算按每个模型的完整继承路径累计，不把两套后缀相加。新方法从第一步使用SR，共20200个SR监督样本；J1和T累计12000，Async2和Sync2累计18000。相同更新数并未配平SR曝光或容量；RGB渲染数也不等同于墙钟耗时。')
    t=s['training'];text('本次训练使用A100 GPU1，训练与保存耗时'+fmt(t['train_seconds'],1)+'秒，峰值分配显存'+fmt(t['peak_allocated_gb'],3)+'GB；完整训练共'+str(t['rgb_forwards'])+'次RGB前向。训练退出前立即完成一次合法训练观测完整性检查。首选SSH链路中断，远端持久训练正常完成；经已有a100-via-5090入口恢复结果回传和本机评价，未重启训练、未使用5090 GPU。实际成本详见complete.json和evaluation日志。统一评价不消除训练硬件与优化轨迹差异。',small=True)

    heading('固定位置图像与时序诊断')
    for camera in ['cam00','cam01']:
        picture(OUT/'figures'/f'{camera}_0040_comparison.png',camera+'第40帧及预先固定区域。HR列为真实图像；历史局部对照为第一套J1和Async2。区域LPIPS来自完整图像上下文特征图，与整图标准LPIPS不同。')
    text('第40和80帧均按既有区域坐标检查。展示图统一量化到PNG，只用于观看；regional_quality.csv的PSNR、SSIM与区域LPIPS来自量化前浮点图。原坐标不是随人体运动追踪的语义区域，局部案例不替代全图主成绩。',small=True)
    text('P1的人脸框PSNR较高，但SSIM未提高；cam01灯具墙角仍有明显失真，其两帧区域PSNR、SSIM与LPIPS均比第一套J1和Async2差。这说明局部像素拟合改善不能替代全场景及多指标判断。',small=True)

    heading('辅助诊断与可追溯记录')
    table(['相机','动态区LPIPS','动态时序误差','整图时序误差','LR闭环L1'],
          [[r['camera'],fmt(r['dynamic_lpips']),fmt(r['dynamic_temporal']),fmt(r['full_temporal']),fmt(r['lr_l1'])] for r in s['auxiliary']],
          [2.5,3.8,3.8,3.8,3.6])
    text('时序误差沿用真实HR的固定光流对应，比较预测帧间变化与真实帧间变化的差；越低越好，但不是几何正确性证明。LR闭环L1比较HR渲染降采样后的图像与实际LR。动态区域来自固定HR时间变化掩码，仅用于评价。')
    tr=s['train16'];text('四个原教师训练相机cam02、cam06、cam12、cam18，在第0、40、80、118帧的固定16观察：PSNR '+fmt(tr['psnr'],4)+' dB，SSIM '+fmt(tr['ssim'])+'，LPIPS '+fmt(tr['lpips'])+'，LR闭环L1 '+fmt(tr['lr_l1'])+'，到SwinIR的RGB平均绝对误差 '+fmt(tr['teacher_rgb_l1'])+'。这是训练拟合诊断，不是新视角性能；未把历史train76平均值当作同分布对照。')
    regional=list(csv.DictReader((OUT/'regional_quality.csv').open()))
    selected=[]
    for camera,region in [('cam00','face_reference'),('cam00','clothing_reference'),('cam01','lamp_wall_corner_reference')]:
        for model in ['r1_J1','r1_Async2','P1-from-start']:
            subset=[r for r in regional if r['camera']==camera and r['region']==region and r['model']==model]
            assert len(subset)==2
            selected.append([camera+' / '+region.replace('_reference','').replace('_',' ')+' / '+model]+
                            [fmt(sum(float(r[k]) for r in subset)/2,4 if k=='psnr' else 6) for k in ['psnr','ssim','lpips_spatial']])
    table(['区域与模型 40和80帧均值','PSNR dB','SSIM','区域LPIPS'],selected,[7.6,3.3,3.3,3.3])
    text('证据源：output/dynamic_sr_same_observation_20260930内的protocol.json、train/config.json、first_update_coarse.json、first_update_fine.json、checkpoint_inventory.json、comparison_quality.csv、quality_gaps.csv、training_budgets.csv、metrics_per_frame.csv、regional_quality.csv及figures/index.json。历史指标只读复用9月29日final_summary.json，来源SHA256已写入comparison_summary.json。',small=True)

    # Keep the separately completed FFT supplement on future report regeneration.
    if (OUT/'frequency/summary.json').exists():
        import importlib.util
        path=ROOT/'experiments/dynamic_sr_same_observation_20260930/frequency_report.py'
        spec=importlib.util.spec_from_file_location('same_observation_frequency_report',path)
        extension=importlib.util.module_from_spec(spec);spec.loader.exec_module(extension)
        extension.append_report(doc,md)


def main():
    backup=OUT/'document_before_update';backup.mkdir(exist_ok=False)
    shutil.copy2(DOC,backup/DOC.name);shutil.copy2(DOC.with_suffix('.md'),backup/DOC.with_suffix('.md').name)
    doc=Document(DOC);md=DOC.with_suffix('.md').read_text().splitlines()
    append_report(doc,md)
    doc.save(DOC);DOC.with_suffix('.md').write_text('\n'.join(md)+'\n')
    print(DOC)


if __name__=='__main__':main()
